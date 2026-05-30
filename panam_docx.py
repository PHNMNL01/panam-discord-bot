import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument


EMPTY_DOCUMENT_TEXT = "Dokument neobsahuje žádný text."
NUMBERED_LIST_RE = re.compile(r"^\s*\d+\.\s+(.+)$")
__all__ = ["create_docx_from_text"]


def _ensure_docx_path(output_path: Path) -> Path:
    if not isinstance(output_path, Path):
        raise TypeError("output_path must be a pathlib.Path object")

    if output_path.suffix.lower() != ".docx":
        return output_path.with_suffix(".docx")

    return output_path


def _strip_inline_markdown(text: str) -> str:
    cleaned = text.strip()
    cleaned = cleaned.strip("`")
    cleaned = re.sub(r"^\*\*(.+)\*\*$", r"\1", cleaned)
    cleaned = re.sub(r"^__(.+)__$", r"\1", cleaned)
    cleaned = re.sub(r"^\*(.+)\*$", r"\1", cleaned)
    cleaned = re.sub(r"^_(.+)_$", r"\1", cleaned)
    return cleaned.strip()


def _add_pending_blank(document: DocxDocument, pending_blank: bool) -> bool:
    if pending_blank:
        document.add_paragraph()
    return False


def create_docx_from_text(
    text: str,
    output_path: Path,
    title: str | None = None,
) -> Path:
    output_path = _ensure_docx_path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()

    if title:
        document.add_heading(_strip_inline_markdown(title), level=1)

    if not text.strip():
        document.add_paragraph(EMPTY_DOCUMENT_TEXT)
        document.save(output_path)
        return output_path

    pending_blank = False

    for raw_line in text.splitlines():
        line = raw_line.strip()

        if not line:
            pending_blank = True
            continue

        pending_blank = _add_pending_blank(document, pending_blank)

        if line.startswith("### "):
            document.add_heading(_strip_inline_markdown(line[4:]), level=3)
            continue

        if line.startswith("## "):
            document.add_heading(_strip_inline_markdown(line[3:]), level=2)
            continue

        if line.startswith("# "):
            document.add_heading(_strip_inline_markdown(line[2:]), level=1)
            continue

        if line.startswith("- ") or line.startswith("* "):
            document.add_paragraph(_strip_inline_markdown(line[2:]), style="List Bullet")
            continue

        numbered_match = NUMBERED_LIST_RE.match(line)
        if numbered_match:
            document.add_paragraph(
                _strip_inline_markdown(numbered_match.group(1)),
                style="List Number",
            )
            continue

        document.add_paragraph(_strip_inline_markdown(line))

    document.save(output_path)
    return output_path
