import json
import re
from pathlib import Path
from typing import Any, cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


MIN_COLUMN_WIDTH = 10
MAX_COLUMN_WIDTH = 50
COLUMN_WIDTH_PADDING = 2


def safe_sheet_name(name: str) -> str:
    sheet_name = re.sub(r"[\[\]:*?/\\]", "_", str(name or "Data")).strip()
    if not sheet_name:
        sheet_name = "Data"
    return sheet_name[:31]


def flatten_value(value: Any) -> str | int | float | bool | None:
    if value is None:
        return ""

    if isinstance(value, (str, int, float, bool)):
        return value

    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)

    return str(value)


def collect_headers(items: list[dict]) -> list[str]:
    headers = []
    seen = set()

    for item in items:
        for key in item.keys():
            header = str(key)
            if header in seen:
                continue

            seen.add(header)
            headers.append(header)

    return headers


def format_worksheet(worksheet: Worksheet) -> None:
    header_font = Font(bold=True)
    top_alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.freeze_panes = "A2"

    for cell in worksheet[1]:
        cell.font = header_font
        cell.alignment = top_alignment

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = top_alignment

    if worksheet.max_row and worksheet.max_column:
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
        )

    for column_index in range(1, worksheet.max_column + 1):
        max_length = 0
        for column_cells in worksheet.iter_cols(
            min_col=column_index,
            max_col=column_index,
            min_row=1,
            max_row=worksheet.max_row,
        ):
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))

        width = min(
            max(max_length + COLUMN_WIDTH_PADDING, MIN_COLUMN_WIDTH),
            MAX_COLUMN_WIDTH,
        )
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def create_xlsx_from_items(
    items: list[dict],
    output_path: Path,
    sheet_name: str = "Data",
) -> Path:
    workbook = Workbook()
    worksheet = cast(Worksheet, workbook.active)
    worksheet.title = safe_sheet_name(sheet_name)

    if not items:
        worksheet.append(["note"])
        worksheet.append(["V dokumentu jsem nenašla požadovaná data."])
    else:
        headers = collect_headers(items)
        if not headers:
            headers = ["value"]

        worksheet.append(headers)
        for item in items:
            worksheet.append([flatten_value(item.get(header, "")) for header in headers])

    format_worksheet(worksheet)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return output_path
