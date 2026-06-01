from io import BytesIO
from pathlib import Path
import sys
import tempfile


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

import panam_text_extraction
from docx import Document
from openpyxl import Workbook


def main() -> None:
    source = Path("panam_text_extraction.py").read_text(encoding="utf-8")
    assert "import discord" not in source
    assert "from discord" not in source

    plain_text = panam_text_extraction.extract_text_from_plain_file(
        "Ahoj Panam".encode("utf-8")
    )
    assert plain_text == "Ahoj Panam", plain_text

    json_text = panam_text_extraction.extract_text_from_attachment(
        "data.json",
        b'{"items":[{"name":"test"}]}',
    )
    assert json_text == '{"items":[{"name":"test"}]}', json_text

    long_text = "x" * (panam_text_extraction.MAX_DOCUMENT_TEXT_LENGTH + 100)
    trimmed_text = panam_text_extraction.trim_document_text(long_text)
    assert len(trimmed_text) <= panam_text_extraction.MAX_DOCUMENT_TEXT_LENGTH
    assert trimmed_text.endswith("...text dokumentu byl zkrácen."), trimmed_text[-80:]

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir)

        docx_path = temp_path / "sample.docx"
        document = Document()
        document.add_paragraph("DOCX text pro Panam")
        document.save(docx_path)

        docx_text = panam_text_extraction.extract_text_from_attachment(
            docx_path.name,
            docx_path.read_bytes(),
        )
        assert "DOCX text pro Panam" in docx_text, docx_text

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Data"
        worksheet.append(["Jméno", "Hodnota"])
        worksheet.append(["Panam", 42])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        xlsx_text = panam_text_extraction.extract_text_from_attachment(
            "sample.xlsx",
            buffer.getvalue(),
        )
        assert "Sheet: Data" in xlsx_text, xlsx_text
        assert "Row 1: Jméno | Hodnota" in xlsx_text, xlsx_text
        assert "Row 2: Panam | 42" in xlsx_text, xlsx_text

    print("text extraction smoke test ok")


if __name__ == "__main__":
    main()
