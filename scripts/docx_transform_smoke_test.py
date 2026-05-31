from pathlib import Path
import sys
from tempfile import TemporaryDirectory

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import panam_docx_transform


def create_input_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("Toto je kratky testovaci dokument.")
    document.add_paragraph("Obsahuje nekolik odstavcu pro upravu.")
    document.add_paragraph("Na konci je poznamka k dalsimu postupu.")
    document.save(path)


def assert_operation(instruction: str, expected_operation: str) -> dict:
    operation = panam_docx_transform.detect_docx_transform_operation(instruction)
    assert operation["operation"] == expected_operation
    assert operation["operation_summary"]
    return operation


def main() -> None:
    runtime_dir = Path(__file__).resolve().parents[1] / "runtime"
    runtime_dir.mkdir(parents=True, exist_ok=True)

    with TemporaryDirectory(dir=runtime_dir) as temp_dir:
        base_dir = Path(temp_dir)
        input_path = base_dir / "dokument.docx"
        create_input_docx(input_path)

        source_text = (
            "Toto je kratky testovaci dokument.\n"
            "Obsahuje nekolik odstavcu pro upravu.\n"
            "Na konci je poznamka k dalsimu postupu."
        )

        test_cases = [
            ("oprava preklepu a stylistiky", "proofread"),
            ("zestrucni text", "summarize"),
            ("preved do formalniho tonu", "formalize"),
            ("vytvor checklist", "checklist"),
        ]
        for index, (instruction, expected_operation) in enumerate(test_cases, start=1):
            operation = assert_operation(instruction, expected_operation)
            output_path = base_dir / f"docx_transform_{index}.docx"
            result = panam_docx_transform.create_docx_transform_from_text(
                source_text,
                f"# Vystup\n\n{source_text}",
                output_path,
                operation["operation_summary"],
            )
            assert output_path.exists(), "Vystupni DOCX neexistuje."
            assert result.operation_summary
            assert result.paragraph_count_before == 3
            assert result.paragraph_count_after >= 3

        try:
            panam_docx_transform.detect_docx_transform_operation(
                "neco tam dopln podle sebe"
            )
        except panam_docx_transform.DocxTransformUserError as error:
            assert str(error) == panam_docx_transform.DOCX_TRANSFORM_FALLBACK_MESSAGE
        else:
            raise AssertionError("Volna DOCX uprava mela skoncit fallbackem.")

    print("docx transform smoke test ok")


if __name__ == "__main__":
    main()
