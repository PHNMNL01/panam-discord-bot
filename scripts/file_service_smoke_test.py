import asyncio
from pathlib import Path
import sys
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from panam_file_service import (
    PanamFileRequest,
    PanamFileResult,
    cleanup_file_result,
    process_panam_file_request,
)


def create_input_xlsx(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Orders"
    worksheet.append(["Name", "Email", "Team"])
    worksheet.append(["Ana Novak", "ana@example.com", "IT"])
    worksheet.append([None, None, None])
    worksheet.append(["Bela Svoboda", "bela@example.com", "HR"])
    workbook.save(path)
    workbook.close()


def read_rows(path: Path) -> list[tuple]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        return list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()


async def exercise_spreadsheet_service(input_path: Path) -> PanamFileResult:
    request = PanamFileRequest(
        input_path=input_path,
        original_filename="orders.xlsx",
        instruction="odstran prazdne radky",
        mode="transform_excel",
        output_format=None,
        user_id="web-user",
        channel_id="1001",
        source="web",
    )
    return await process_panam_file_request("", request)


def main() -> None:
    request = PanamFileRequest(
        input_path=Path("input.xlsx"),
        original_filename="input.xlsx",
        instruction="odstran prazdne radky",
        mode="spreadsheet_transform",
        output_format="xlsx",
        user_id=1,
        channel_id=2,
        source="discord",
    )
    assert request.original_filename == "input.xlsx"

    result = PanamFileResult(
        kind="chat_answer",
        message="ok",
        output_path=None,
        output_filename=None,
        mode="chat_answer",
        output_format=None,
    )
    assert result.kind == "chat_answer"
    assert result.output_path is None

    with TemporaryDirectory() as temp_dir:
        input_path = Path(temp_dir) / "orders.xlsx"
        create_input_xlsx(input_path)
        original_rows = read_rows(input_path)

        service_result = asyncio.run(exercise_spreadsheet_service(input_path))
        try:
            assert service_result.kind == "file"
            assert service_result.mode == "spreadsheet_transform"
            assert service_result.output_format == "xlsx"
            assert service_result.output_filename == "orders_by_Panam.xlsx"
            assert service_result.output_path is not None
            assert service_result.output_path.exists()
            assert service_result.row_count_before == 3
            assert service_result.row_count_after == 2
            assert read_rows(input_path) == original_rows

            output_rows = read_rows(service_result.output_path)
            assert output_rows == [
                ("Name", "Email", "Team"),
                ("Ana Novak", "ana@example.com", "IT"),
                ("Bela Svoboda", "bela@example.com", "HR"),
            ]
        finally:
            cleanup_file_result(service_result)

    print("file service smoke test ok")


if __name__ == "__main__":
    main()
