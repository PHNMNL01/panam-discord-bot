from pathlib import Path
import sys
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import panam_spreadsheet


def create_input_xlsx(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Export"
    worksheet.append(
        [
            "Jméno a příjmení",
            "Soukromý e-mail",
            "Název pracovní pozice",
            "Oddělení",
            "Datum nástupu",
        ]
    )
    worksheet.append(["Ana Novak", "ana@example.com", "Analytik", "IT", "2024-01-15"])
    worksheet.append([None, None, None, None, None])
    worksheet.append(["Bela Svoboda", "bela@example.com", "HR specialista", "HR", "2023-06-01"])
    worksheet.append(["Cyril Dvorak", "cyril@example.com", "Vyvojar", "IT", "2022-03-20"])
    worksheet.append(["Daria Novak", "ana@example.com", "Tester", "HR", "2025-02-10"])
    workbook.save(path)
    workbook.close()


def create_titled_header_input_xlsx(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Vozovy park"
    worksheet.append(["VOZOVÝ PARK - interní export"])
    worksheet.append([])
    worksheet.append([])
    worksheet.append(["SPZ", "ZNAČKA", "ŘIDIČ"])
    worksheet.append(["1AB 2345", "ŠKODA RAPID", "Ana Novak"])
    worksheet.append(["2CD 6789", "VW GOLF", "Bela Svoboda"])
    worksheet.append(["3EF 0123", "ŠKODA RAPID", "Cyril Dvorak"])
    workbook.save(path)
    workbook.close()


def read_rows(path: Path) -> list[tuple]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        return list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def main() -> None:
    runtime_dir = Path(__file__).resolve().parents[1] / "runtime"
    runtime_dir.mkdir(parents=True, exist_ok=True)

    with TemporaryDirectory(dir=runtime_dir) as temp_dir:
        base_dir = Path(temp_dir)
        input_path = base_dir / "export.xlsx"
        create_input_xlsx(input_path)

        output_path = base_dir / "export_remove_empty_by_Panam.xlsx"
        result = panam_spreadsheet.transform_xlsx(
            input_path,
            output_path,
            "odstran prazdne radky",
        )
        assert output_path.exists(), "Vystupni XLSX neexistuje."
        assert result.operation_summary == "odstranění prázdných řádků"
        assert result.row_count_before == 5
        assert result.row_count_after == 4

        output_path = base_dir / "export_select_email_by_Panam.xlsx"
        result = panam_spreadsheet.transform_xlsx(
            input_path,
            output_path,
            "vyber sloupce Jmeno a Prijmeni, Soukromy email",
        )
        assert output_path.exists(), "Vystupni XLSX neexistuje."
        assert result.operation_summary == "výběr sloupců Jméno a příjmení, Soukromý e-mail"
        assert result.column_count_before == 5
        assert result.column_count_after == 2
        rows = read_rows(output_path)
        assert rows[0] == ("Jméno a příjmení", "Soukromý e-mail")

        output_path = base_dir / "export_select_position_by_Panam.xlsx"
        result = panam_spreadsheet.transform_xlsx(
            input_path,
            output_path,
            "vyber sloupce Jmeno a prijmeni, Nazev pracovni pozice",
        )
        assert output_path.exists(), "Vystupni XLSX neexistuje."
        assert result.operation_summary == "výběr sloupců Jméno a příjmení, Název pracovní pozice"
        rows = read_rows(output_path)
        assert rows[0] == ("Jméno a příjmení", "Název pracovní pozice")

        output_path = base_dir / "export_filter_it_by_Panam.xlsx"
        result = panam_spreadsheet.transform_xlsx(
            input_path,
            output_path,
            "nech jen řádky kde Oddeleni = IT",
        )
        assert output_path.exists(), "Vystupni XLSX neexistuje."
        assert result.operation_summary == "filtr řádků, kde Oddělení = IT"
        assert result.row_count_before == 5
        assert result.row_count_after == 2
        rows = read_rows(output_path)
        assert rows[0] == (
            "Jméno a příjmení",
            "Soukromý e-mail",
            "Název pracovní pozice",
            "Oddělení",
            "Datum nástupu",
        )
        assert len(rows) == 3

        output_path = base_dir / "export_sort_date_by_Panam.xlsx"
        result = panam_spreadsheet.transform_xlsx(
            input_path,
            output_path,
            "seřaď podle Datum nastupu",
        )
        assert output_path.exists(), "Vystupni XLSX neexistuje."
        assert result.operation_summary == "seřazení podle sloupce Datum nástupu"
        rows = read_rows(output_path)
        assert rows[0] == (
            "Jméno a příjmení",
            "Soukromý e-mail",
            "Název pracovní pozice",
            "Oddělení",
            "Datum nástupu",
        )

        output_path = base_dir / "export_duplicates_by_Panam.xlsx"
        result = panam_spreadsheet.transform_xlsx(
            input_path,
            output_path,
            "najdi duplicity podle Soukromy email",
        )
        assert output_path.exists(), "Vystupni XLSX neexistuje."
        assert result.operation_summary == "nalezení duplicit podle sloupce Soukromý e-mail"

        titled_input_path = base_dir / "vozovy_park.xlsx"
        create_titled_header_input_xlsx(titled_input_path)

        output_path = base_dir / "vozovy_park_skoda_rapid_by_Panam.xlsx"
        result = panam_spreadsheet.transform_xlsx(
            titled_input_path,
            output_path,
            "nech jen řádky kde Značka = ŠKODA RAPID",
        )
        assert output_path.exists(), "Vystupni XLSX neexistuje."
        assert result.row_count_before == 3
        assert result.row_count_after == 2
        rows = read_rows(output_path)
        assert rows[0] == ("SPZ", "ZNAČKA", "ŘIDIČ")
        assert len(rows) == 3
        assert all(row[1] == "ŠKODA RAPID" for row in rows[1:])

    print("spreadsheet transform smoke test ok")


if __name__ == "__main__":
    main()
