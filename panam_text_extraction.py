import io
import logging
from pathlib import Path


MAX_DOCUMENT_TEXT_LENGTH = 20000
MAX_XLSX_SHEETS = 10
MAX_XLSX_ROWS_PER_SHEET = 500
MAX_XLSX_COLUMNS_PER_SHEET = 50

logger = logging.getLogger("panam")


class TextExtractionUserError(Exception):
    pass


def get_file_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def extract_text_from_plain_file(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("utf-8", errors="replace")


def extract_text_from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    page_texts = []

    for page in reader.pages:
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            page_texts.append(text)

    return "\n\n".join(page_texts).strip()


def extract_text_from_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cell_texts = []
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    cell_texts.append(text)

            if cell_texts:
                lines.append(" | ".join(cell_texts))

    return "\n".join(lines).strip()


def extract_text_from_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(
            io.BytesIO(data),
            data_only=True,
            read_only=True,
        )
    except Exception as error:
        logger.info(
            "xlsx_extraction status=error reason=open_failed size_bytes=%s",
            len(data),
        )
        raise TextExtractionUserError(
            "Ten Excel se mi nepodařilo otevřít. Soubor může být poškozený nebo v nepodporovaném formátu."
        ) from error

    lines = []
    sheet_count = 0
    used_rows = 0
    used_cols = 0

    try:
        for worksheet in workbook.worksheets[:MAX_XLSX_SHEETS]:
            sheet_lines = []
            sheet_used_rows = 0
            sheet_used_cols = 0

            for row_index, row in enumerate(
                worksheet.iter_rows(
                    max_row=MAX_XLSX_ROWS_PER_SHEET,
                    max_col=MAX_XLSX_COLUMNS_PER_SHEET,
                    values_only=True,
                ),
                start=1,
            ):
                row_values = ["" if cell is None else str(cell).strip() for cell in row]
                non_empty_indexes = [
                    index
                    for index, value in enumerate(row_values)
                    if value
                ]
                if not non_empty_indexes:
                    continue

                last_value_index = max(non_empty_indexes)
                trimmed_values = row_values[: last_value_index + 1]
                sheet_lines.append(
                    f"Row {row_index}: " + " | ".join(trimmed_values)
                )
                sheet_used_rows += 1
                sheet_used_cols = max(sheet_used_cols, last_value_index + 1)

            if sheet_lines:
                sheet_count += 1
                used_rows += sheet_used_rows
                used_cols = max(used_cols, sheet_used_cols)
                if lines:
                    lines.append("")
                lines.append(f"Sheet: {worksheet.title}")
                lines.extend(sheet_lines)

    finally:
        workbook.close()

    extracted_text = "\n".join(lines).strip()
    logger.info(
        "xlsx_extraction status=%s sheet_count=%s used_rows=%s used_cols=%s",
        "success" if extracted_text else "empty",
        sheet_count,
        used_rows,
        used_cols,
    )
    return trim_document_text(extracted_text)


def extract_text_from_attachment(filename: str, data: bytes) -> str:
    extension = get_file_extension(filename)
    if extension == ".pdf":
        return extract_text_from_pdf(data)
    if extension == ".docx":
        return extract_text_from_docx(data)
    if extension == ".xlsx":
        return extract_text_from_xlsx(data)
    return extract_text_from_plain_file(data)


def trim_document_text(text: str) -> str:
    if len(text) <= MAX_DOCUMENT_TEXT_LENGTH:
        return text

    suffix = "\n\n...text dokumentu byl zkrácen."
    return text[: max(MAX_DOCUMENT_TEXT_LENGTH - len(suffix), 0)].rstrip() + suffix
