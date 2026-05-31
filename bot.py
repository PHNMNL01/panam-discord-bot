import os
import asyncio
import io
import logging
import logging.handlers
import json
import re
import unicodedata
from functools import wraps
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import discord
from discord import app_commands
from dotenv import load_dotenv

import panam_docx
import panam_excel
import panam_files
import panam_memory
import panam_spreadsheet
from panam_file_context import (
    clear_last_file_context,
    clear_last_router_decision,
    get_last_file_context,
    set_last_file_context,
    set_last_router_decision,
)
from panam_ai import (
    analyze_image,
    analyze_document_text,
    ask_panam,
    ask_panam_talk,
    classify_file_request_intent,
    extract_structured_data,
    process_document_text,
    shorten_for_discord,
    summarize_channel_messages,
    summarize_text,
)
from panam_phrases import (
    ATTACHMENT_ANALYZE_PATTERNS,
    ATTACHMENT_SUBJECTS,
    BASIC_PANAM_EMPTY_RESPONSE,
    CONTEXT_REFERENCES,
    GENERIC_ATTACHMENT_PATTERNS,
    GENERIC_IMAGE_PHRASES,
    HELP_PATTERNS,
    IMAGE_ANALYZE_TRIGGERS,
    NATURAL_INTENT_PATTERNS,
    NOTE_ADD_PREVIOUS_PATTERN,
    NOTE_ADD_TRIGGERS,
    NOTE_LIST_PATTERN,
    OPINION_CONTEXT_PATTERNS,
    OPINION_TRIGGERS,
    PANAM_OPINION_MENTION_PATTERN,
    PANAM_PREFIX_PATTERN,
    PANAM_STRIP_PREFIX_PATTERN,
    SUMMARY_CONTEXT_PATTERN,
    SUMMARY_TRIGGERS,
    TODO_LIST_PATTERN,
)
from panam_router import (
    decide_file_response_mode,
    has_direct_file_edit_request,
    has_explicit_file_action_request,
    has_explicit_file_output_request,
    has_spreadsheet_transform_request,
    has_unsafe_creative_spreadsheet_request,
    is_ambiguous_context_request,
    is_followup_to_file_summary,
    is_meta_router_or_behavior_discussion,
    matches_last_file_reference,
    validate_ai_file_intent,
)


BASE_DIR = Path(__file__).resolve().parent
LOG_DIR = BASE_DIR / "logs"
LOG_FILE = LOG_DIR / "panam.log"

load_dotenv(dotenv_path=BASE_DIR / ".env")

DISCORD_BOT_TOKEN = os.getenv("DISCORD_BOT_TOKEN")
DISCORD_GUILD_IDS = [
    guild_id.strip()
    for guild_id in os.getenv("DISCORD_GUILD_IDS", "").split(",")
    if guild_id.strip()
]
ALLOWED_CHANNEL_IDS = [
    channel_id.strip()
    for channel_id in os.getenv("ALLOWED_CHANNEL_IDS", "").split(",")
    if channel_id.strip()
]
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4.1-mini")
NOTES_FILE = BASE_DIR / "notes.json"
TODOS_FILE = BASE_DIR / "todos.json"
SUPPORTED_IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".webp", ".gif")
SUPPORTED_DOCUMENT_EXTENSIONS = (".txt", ".md", ".csv", ".pdf", ".docx", ".xlsx")
MAX_IMAGE_SIZE_BYTES = 20 * 1024 * 1024
MAX_DOCUMENT_SIZE_BYTES = 20 * 1024 * 1024
MAX_DOCUMENT_TEXT_LENGTH = 20000
MAX_XLSX_SHEETS = 10
MAX_XLSX_ROWS_PER_SHEET = 500
MAX_XLSX_COLUMNS_PER_SHEET = 50
COMMAND_STATUSES: dict[int, str] = {}
CREATIVE_SPREADSHEET_FALLBACK_MESSAGE = (
    "Tohle je moc volná úprava. Původní Excel neupravuju a data si nedomýšlím. "
    "Umím bezpečně udělat nový XLSX třeba: odstranit prázdné řádky, filtrovat řádky podle sloupce, "
    "vybrat sloupce, seřadit podle sloupce nebo najít duplicity."
)


logger = logging.getLogger("panam")


def setup_logging() -> None:
    LOG_DIR.mkdir(parents=True, exist_ok=True)

    formatter = logging.Formatter("%(asctime)s [%(levelname)s] %(name)s: %(message)s")
    root_logger = logging.getLogger()
    root_logger.setLevel(logging.INFO)

    if not any(isinstance(handler, logging.StreamHandler) for handler in root_logger.handlers):
        console_handler = logging.StreamHandler()
        console_handler.setFormatter(formatter)
        root_logger.addHandler(console_handler)

    if not any(
        isinstance(handler, logging.handlers.RotatingFileHandler)
        and Path(handler.baseFilename) == LOG_FILE
        for handler in root_logger.handlers
    ):
        file_handler = logging.handlers.RotatingFileHandler(
            LOG_FILE,
            maxBytes=1024 * 1024,
            backupCount=5,
            encoding="utf-8",
        )
        file_handler.setFormatter(formatter)
        root_logger.addHandler(file_handler)


def get_safe_user_name(user) -> str:
    return str(getattr(user, "display_name", getattr(user, "name", "unknown")))


def get_safe_context(source) -> dict[str, str | int | None]:
    user = getattr(source, "user", None) or getattr(source, "author", None)
    channel_id = getattr(source, "channel_id", None)
    if channel_id is None:
        channel = getattr(source, "channel", None)
        channel_id = getattr(channel, "id", None)

    guild_id = getattr(source, "guild_id", None)
    if guild_id is None:
        guild = getattr(source, "guild", None)
        guild_id = getattr(guild, "id", None)

    return {
        "user_id": getattr(user, "id", None),
        "user_name": get_safe_user_name(user) if user is not None else "unknown",
        "channel_id": channel_id,
        "guild_id": guild_id,
    }


def get_context_int(context: dict[str, str | int | None], key: str) -> int:
    value = context.get(key)
    if isinstance(value, int):
        return value
    if isinstance(value, str) and value.isdecimal():
        return int(value)
    return 0


def get_safe_attachment_info(file: discord.Attachment | None) -> dict[str, str | int | None]:
    if file is None:
        return {}

    extension = get_file_extension(file.filename)
    return {
        "filename": Path(file.filename).name,
        "extension": extension,
        "size": file.size,
        "file_type": get_attachment_kind(file),
    }


def log_attachment_info(action_type: str, action_name: str, source, file: discord.Attachment) -> None:
    context = get_safe_context(source)
    attachment = get_safe_attachment_info(file)
    logger.info(
        "action_type=%s action=%s user_id=%s user_name=%s channel_id=%s guild_id=%s filename=%s extension=%s size=%s file_type=%s",
        action_type,
        action_name,
        context["user_id"],
        context["user_name"],
        context["channel_id"],
        context["guild_id"],
        attachment.get("filename"),
        attachment.get("extension"),
        attachment.get("size"),
        attachment.get("file_type"),
    )


def log_action(
    action_type: str,
    action_name: str,
    status: str,
    source,
    **details,
) -> None:
    context = get_safe_context(source)
    detail_text = " ".join(
        f"{key}={value}"
        for key, value in details.items()
        if value is not None
    )
    logger.info(
        "action_type=%s action=%s status=%s user_id=%s user_name=%s channel_id=%s guild_id=%s%s",
        action_type,
        action_name,
        status,
        context["user_id"],
        context["user_name"],
        context["channel_id"],
        context["guild_id"],
        f" {detail_text}" if detail_text else "",
    )


def mark_command_status(interaction: discord.Interaction, status: str) -> None:
    COMMAND_STATUSES[interaction.id] = status


def is_interaction_allowed(interaction: discord.Interaction, command_name: str) -> bool:
    if ALLOWED_CHANNEL_IDS and str(interaction.channel_id) not in ALLOWED_CHANNEL_IDS:
        log_action("slash_command", command_name, "denied", interaction)
        mark_command_status(interaction, "denied")
        return False

    return True


def log_slash_command(command_name: str):
    def decorator(func):
        @wraps(func)
        async def wrapper(interaction: discord.Interaction, *args, **kwargs):
            log_action("slash_command", command_name, "started", interaction)
            try:
                result = await func(interaction, *args, **kwargs)
            except Exception:
                COMMAND_STATUSES.pop(interaction.id, None)
                log_action("slash_command", command_name, "error", interaction)
                logger.exception("Neosetrena chyba pri zpracovani /%s", command_name)
                raise

            status = COMMAND_STATUSES.pop(interaction.id, None)
            if status is None:
                log_action("slash_command", command_name, "success", interaction)
            elif status == "error":
                log_action("slash_command", command_name, "error", interaction)

            return result

        return wrapper

    return decorator


setup_logging()
logger.info("Panam bot startuje.")


if not DISCORD_BOT_TOKEN:
    raise RuntimeError("Chybí DISCORD_BOT_TOKEN v .env souboru.")


def load_notes() -> list[dict]:
    if not NOTES_FILE.exists():
        return []

    with NOTES_FILE.open("r", encoding="utf-8") as notes_file:
        notes = json.load(notes_file)

    if not isinstance(notes, list):
        return []

    return notes


def save_notes(notes: list[dict]) -> None:
    with NOTES_FILE.open("w", encoding="utf-8") as notes_file:
        json.dump(notes, notes_file, ensure_ascii=False, indent=2)


def load_todos() -> list[dict]:
    if not TODOS_FILE.exists():
        return []

    with TODOS_FILE.open("r", encoding="utf-8") as todos_file:
        todos = json.load(todos_file)

    if not isinstance(todos, list):
        return []

    return todos


def save_todos(todos: list[dict]) -> None:
    with TODOS_FILE.open("w", encoding="utf-8") as todos_file:
        json.dump(todos, todos_file, ensure_ascii=False, indent=2)


def get_author_name(author) -> str:
    return getattr(author, "display_name", author.name)


def split_discord_message(text: str, limit: int = 1900) -> list[str]:
    if not text:
        return [""]

    chunks = []
    remaining = text.strip()

    while len(remaining) > limit:
        split_at = remaining.rfind("\n", 0, limit + 1)
        if split_at <= 0:
            split_at = remaining.rfind(" ", 0, limit + 1)
        if split_at <= 0:
            split_at = limit

        chunks.append(remaining[:split_at].rstrip())
        remaining = remaining[split_at:].lstrip()

    if remaining:
        chunks.append(remaining)

    return chunks


def get_file_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def is_supported_image_attachment(attachment: discord.Attachment) -> bool:
    return get_file_extension(attachment.filename) in SUPPORTED_IMAGE_EXTENSIONS


def is_supported_document_attachment(attachment: discord.Attachment) -> bool:
    return get_file_extension(attachment.filename) in SUPPORTED_DOCUMENT_EXTENSIONS


def get_attachment_kind(file: discord.Attachment) -> str | None:
    extension = get_file_extension(file.filename)
    if extension in SUPPORTED_IMAGE_EXTENSIONS:
        return "image"
    if extension in SUPPORTED_DOCUMENT_EXTENSIONS:
        return "document"
    return None


class AttachmentAnalysisUserError(Exception):
    pass


def find_image_attachment_in_message(
    message: discord.Message,
) -> discord.Attachment | None:
    for attachment in message.attachments:
        if is_supported_image_attachment(attachment):
            return attachment

    return None


def find_supported_attachment_in_message(
    message: discord.Message,
) -> discord.Attachment | None:
    for attachment in message.attachments:
        if get_attachment_kind(attachment) is not None:
            return attachment

    return None


async def read_attachment_bytes(file: discord.Attachment) -> bytes:
    return await file.read()


def extract_text_from_plain_file(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("utf-8", errors="replace")


def extract_text_from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    page_texts = []

    for page in reader.pages:
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            page_texts.append(text)

    return "\n\n".join(page_texts).strip()


def extract_text_from_attachment(filename: str, data: bytes) -> str:
    extension = get_file_extension(filename)
    if extension == ".pdf":
        return extract_text_from_pdf(data)
    if extension == ".docx":
        return extract_text_from_docx(data)
    if extension == ".xlsx":
        return extract_text_from_xlsx(data)
    return extract_text_from_plain_file(data)


def extract_text_from_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cell_texts = []
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    cell_texts.append(text)

            if cell_texts:
                lines.append(" | ".join(cell_texts))

    return "\n".join(lines).strip()


def extract_text_from_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(
            io.BytesIO(data),
            data_only=True,
            read_only=True,
        )
    except Exception as error:
        logger.info(
            "xlsx_extraction status=error reason=open_failed size_bytes=%s",
            len(data),
        )
        raise AttachmentAnalysisUserError(
            "Ten Excel se mi nepodařilo otevřít. Soubor může být poškozený nebo v nepodporovaném formátu."
        ) from error

    lines = []
    sheet_count = 0
    used_rows = 0
    used_cols = 0

    try:
        for worksheet in workbook.worksheets[:MAX_XLSX_SHEETS]:
            sheet_lines = []
            sheet_used_rows = 0
            sheet_used_cols = 0

            for row_index, row in enumerate(
                worksheet.iter_rows(
                    max_row=MAX_XLSX_ROWS_PER_SHEET,
                    max_col=MAX_XLSX_COLUMNS_PER_SHEET,
                    values_only=True,
                ),
                start=1,
            ):
                row_values = ["" if cell is None else str(cell).strip() for cell in row]
                non_empty_indexes = [
                    index
                    for index, value in enumerate(row_values)
                    if value
                ]
                if not non_empty_indexes:
                    continue

                last_value_index = max(non_empty_indexes)
                trimmed_values = row_values[: last_value_index + 1]
                sheet_lines.append(
                    f"Row {row_index}: " + " | ".join(trimmed_values)
                )
                sheet_used_rows += 1
                sheet_used_cols = max(sheet_used_cols, last_value_index + 1)

            if sheet_lines:
                sheet_count += 1
                used_rows += sheet_used_rows
                used_cols = max(used_cols, sheet_used_cols)
                if lines:
                    lines.append("")
                lines.append(f"Sheet: {worksheet.title}")
                lines.extend(sheet_lines)

    finally:
        workbook.close()

    extracted_text = "\n".join(lines).strip()
    logger.info(
        "xlsx_extraction status=%s sheet_count=%s used_rows=%s used_cols=%s",
        "success" if extracted_text else "empty",
        sheet_count,
        used_rows,
        used_cols,
    )
    return trim_document_text(extracted_text)


def trim_document_text(text: str) -> str:
    if len(text) <= MAX_DOCUMENT_TEXT_LENGTH:
        return text

    suffix = "\n\n...text dokumentu byl zkrácen."
    return text[: max(MAX_DOCUMENT_TEXT_LENGTH - len(suffix), 0)].rstrip() + suffix


async def find_recent_image_attachment(channel) -> discord.Attachment | None:
    history = getattr(channel, "history", None)
    if history is None:
        return None

    async for message in history(limit=15):
        if message.author.bot:
            continue

        for attachment in message.attachments:
            if is_supported_image_attachment(attachment):
                return attachment

    return None


async def find_recent_supported_attachment(channel) -> discord.Attachment | None:
    history = getattr(channel, "history", None)
    if history is None:
        return None

    async for message in history(limit=15):
        if message.author.bot:
            continue

        for attachment in message.attachments:
            if get_attachment_kind(attachment) is not None:
                return attachment

    return None


async def find_recent_xlsx_attachment(channel) -> discord.Attachment | None:
    history = getattr(channel, "history", None)
    if history is None:
        return None

    async for message in history(limit=15):
        if message.author.bot:
            continue

        for attachment in message.attachments:
            if get_file_extension(attachment.filename) == ".xlsx":
                return attachment

    return None


async def analyze_selected_attachment(file: discord.Attachment, question: str) -> str:
    attachment_kind = get_attachment_kind(file)
    if attachment_kind is None:
        raise AttachmentAnalysisUserError(
            "Tenhle typ souboru zatím neumím přečíst. Podporuju obrázky PNG, JPG, JPEG, WEBP, GIF a dokumenty TXT, MD, CSV, PDF, DOCX, XLSX."
        )

    if file.size > MAX_IMAGE_SIZE_BYTES:
        raise AttachmentAnalysisUserError(
            "Ten soubor je moc velký. Zatím beru max 20 MB."
        )

    if attachment_kind == "image":
        answer = await analyze_image(OPENAI_MODEL, file.url, question)
        return shorten_for_discord(answer)

    data = await read_attachment_bytes(file)
    document_text = extract_text_from_attachment(file.filename, data)

    if not document_text.strip():
        if get_file_extension(file.filename) == ".xlsx":
            raise AttachmentAnalysisUserError(
                "Z toho Excelu se mi nepodarilo vytahnout zadna data."
            )

        if get_file_extension(file.filename) == ".pdf":
            raise AttachmentAnalysisUserError(
                "Z toho PDF se mi nepodařilo vytáhnout žádný text. Možná je to sken nebo obrázkové PDF."
            )

        raise AttachmentAnalysisUserError(
            "Z toho dokumentu se mi nepodařilo vytáhnout žádný text."
        )

    document_text = trim_document_text(document_text)
    answer = await analyze_document_text(
        OPENAI_MODEL,
        document_text,
        question,
        file.filename,
    )
    return shorten_for_discord(answer)


async def send_followup_chunks(interaction: discord.Interaction, text: str) -> None:
    for chunk in split_discord_message(text):
        await interaction.followup.send(chunk)


async def send_channel_chunks(message: discord.Message, text: str) -> None:
    for chunk in split_discord_message(text):
        await message.channel.send(chunk)


def create_note(text: str, author, channel_id: int) -> None:
    notes = load_notes()
    notes.append(
        {
            "text": text,
            "author_id": author.id,
            "author_name": get_author_name(author),
            "channel_id": channel_id,
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
    )
    save_notes(notes)


def format_note_list_response() -> str:
    notes = load_notes()
    if not notes:
        return "Zatím nemám žádné poznámky."

    lines = ["Poslední poznámky:"]
    for index, note in enumerate(reversed(notes[-10:]), start=1):
        author_name = note.get("author_name", "neznámý autor")
        created_at = note.get("created_at", "neznámý čas")
        note_text = note.get("text", "")
        lines.append(f"{index}. [{created_at}] {author_name}: {note_text}")

    return shorten_for_discord("\n".join(lines))


def format_note_search_response(query: str) -> str:
    notes = load_notes()
    if not notes:
        return "Zatím nemám žádné poznámky."

    query_lower = query.lower()
    matches = [
        note
        for note in notes
        if query_lower in str(note.get("text", "")).lower()
    ][-10:]

    if not matches:
        return "Nic jsem nenašla."

    lines = ["Nalezené poznámky:"]
    for index, note in enumerate(reversed(matches), start=1):
        author_name = note.get("author_name", "neznámý autor")
        created_at = note.get("created_at", "neznámý čas")
        note_text = note.get("text", "")
        lines.append(f"{index}. [{created_at}] {author_name}: {note_text}")

    return shorten_for_discord("\n".join(lines))


def create_todo(text: str, author, channel_id: int) -> int:
    todos = load_todos()
    next_id = max(
        (todo.get("id", 0) for todo in todos if isinstance(todo.get("id"), int)),
        default=0,
    ) + 1

    todos.append(
        {
            "id": next_id,
            "text": text,
            "done": False,
            "author_id": author.id,
            "author_name": get_author_name(author),
            "channel_id": channel_id,
            "created_at": datetime.now(timezone.utc).isoformat(),
            "completed_at": None,
        }
    )
    save_todos(todos)

    return next_id


def format_todo_list_response() -> str:
    todos = load_todos()
    active_todos = [todo for todo in todos if not todo.get("done")]
    if not active_todos:
        return "Nemáš žádné aktivní úkoly."

    lines = ["Aktivní úkoly:"]
    for todo in active_todos[:15]:
        todo_id = todo.get("id", "?")
        text = todo.get("text", "")
        author_name = todo.get("author_name", "neznámý autor")
        lines.append(f"#{todo_id} - {text} ({author_name})")

    return shorten_for_discord("\n".join(lines))


def get_help_text() -> str:
    return (
        "Panam nápověda\n\n"
        "1. Slash commandy\n"
        "`/ask`, `/summary`, `/channel_summary`, `/search_messages`, `/panam_talk`\n"
        "`/note_add`, `/note_list`, `/note_search`, `/todo_add`, `/todo_list`, `/todo_done`\n"
        "`/analyze`, `/read_file`, `/process_file`, `/extract_data`, `/file_job_test`\n"
        "`/memory_clear`, `/ping`, `/help`\n\n"
        "2. Běžný chat\n"
        "`Panam <dotaz>`, `Panam řekni mi <dotaz>`, `Panam odpověz <dotaz>`\n"
        "`Panam co si myslíš o <text>`, `Panam co si o tom myslíš?`\n"
        "`Panam shrň <text>`, `Panam shrň to`, `Panam vysvětli to`\n\n"
        "3. Poznámky a todo\n"
        "`Panam přidej poznámku <text>`, `Panam zapamatuj si to`, "
        "`Panam ukaž poznámky`, `Panam najdi poznámku <text>`\n"
        "`Panam přidej todo <text>`, `Panam přidej úkol <text>`, "
        "`Panam ukaž todo`, `Panam ukaž úkoly`\n\n"
        "4. Přílohy a soubory\n"
        "Obrázky: `.png`, `.jpg`, `.jpeg`, `.webp`, `.gif`. "
        "Dokumenty: `.txt`, `.md`, `.csv`, `.pdf`, `.docx`, `.xlsx` do 20 MB.\n"
        "`/analyze` nebo `Panam co je v tom souboru?` odpoví do chatu.\n"
        "`/read_file` přečte dokument. `/process_file` vytvoří `.md`/`.txt`/`.docx`.\n"
        "`/extract_data` vytěží `json`, `csv`, `md` nebo `xlsx`.\n\n"
        "5. Přirozené file požadavky\n"
        "`Panam shrň ten soubor`, `Panam najdi chyby v té tabulce` -> odpověď do chatu\n"
        "`Panam udělej z toho report`, `Panam dej mi to do souboru` -> nový MD/TXT/DOCX soubor\n"
        "`Panam dej mi to do Excelu`, `Panam requirements do csv`, "
        "`Panam udělej z requirements tabulku` -> strukturovaný výstup\n"
        "Název posledního souboru, třeba `requirements`, umím spojit s file contextem. "
        "Obecné `co je na tom špatně?` nechávám chatu, pokud není jasné, že jde o soubor.\n\n"
        "6. Krátká RAM paměť\n"
        "Pamatuju si krátkou historii, poslední file context, bezpečné shrnutí souboru "
        "a poslední router rozhodnutí. Jen v RAM, po restartu pryč.\n"
        "`/memory_clear` smaže paměť pro aktuální kanál.\n\n"
        "7. Talk mód\n"
        "`/panam_talk <text>`, `Panam talk <text>`, `Panam pokec <text>`, "
        "`Panam co si fakt myslíš o <text>`\n\n"
        "8. Bezpečnost\n"
        "Původní přílohu neupravuju, vytvářím nový XLSX, CSV, MD nebo TXT. "
        "Neposílej hesla, tokeny, API klíče ani citlivá data."
    )


async def find_previous_message_content(
    message: discord.Message,
    prefer_same_author: bool = True,
) -> Optional[str]:
    history = getattr(message.channel, "history", None)
    if history is None:
        return None

    fallback_content = None

    async for previous_message in history(limit=30, before=message.created_at):
        if previous_message.author.bot:
            continue

        previous_content = (previous_message.content or "").strip()
        if not previous_content:
            continue

        if not prefer_same_author:
            return previous_content

        if previous_message.author.id == message.author.id:
            return previous_content

        if fallback_content is None:
            fallback_content = previous_content

    return fallback_content


def extract_panam_request(
    message: discord.Message,
    bot_user: discord.ClientUser,
) -> Optional[str]:
    content = (message.content or "").strip()
    mention_patterns = (
        f"<@{bot_user.id}>",
        f"<@!{bot_user.id}>",
    )

    for mention in mention_patterns:
        if content.startswith(mention):
            return content[len(mention):].strip(" \t\n\r,.:;!-")

    match = re.match(PANAM_PREFIX_PATTERN, content, re.IGNORECASE)
    if match:
        return match.group(1).strip()

    if re.match(
        PANAM_OPINION_MENTION_PATTERN,
        content,
        re.IGNORECASE,
    ):
        return content

    return None


def extract_basic_panam_prompt(content: str) -> str | None:
    match = re.match(PANAM_PREFIX_PATTERN, content.strip(), re.IGNORECASE)
    if not match:
        return None

    return match.group(1).strip()


async def handle_basic_panam_message(
    message: discord.Message,
    basic_prompt: str,
) -> None:
    if not basic_prompt.strip():
        answer = BASIC_PANAM_EMPTY_RESPONSE
        await message.reply(answer, mention_author=False)
        panam_memory.add_message(message.channel.id, "user", message.content or "Panam")
        panam_memory.add_message(message.channel.id, "assistant", answer)
        return

    history = panam_memory.get_messages(message.channel.id)
    answer = await ask_panam(OPENAI_MODEL, basic_prompt, history=history)
    await message.reply(answer, mention_author=False)
    panam_memory.add_message(message.channel.id, "user", message.content or basic_prompt)
    panam_memory.add_message(message.channel.id, "assistant", answer)


def normalize_natural_text(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", text.lower())
    without_diacritics = "".join(
        character
        for character in normalized
        if not unicodedata.combining(character)
    )
    return re.sub(r"\s+", " ", without_diacritics).strip(" \t\n\r,.:;!-?")


def normalize_text(text: str) -> str:
    return normalize_natural_text(text)


def is_panam_addressed(content: str) -> bool:
    text = normalize_text(content)
    return (
        text.startswith("panam")
        or text.startswith("hey panam")
        or re.search(r"\bpanam\b", text) is not None
    )


def is_context_reference(text: str) -> bool:
    normalized = normalize_text(text)
    return normalized in CONTEXT_REFERENCES


def extract_inline_content_after_trigger(content: str, trigger_phrases: list[str]) -> str:
    text = content.strip()
    text = re.sub(PANAM_STRIP_PREFIX_PATTERN, "", text, flags=re.IGNORECASE)

    for phrase in trigger_phrases:
        pattern = rf"^{re.escape(phrase)}\b\s*(.*)$"
        match = re.match(pattern, text, re.IGNORECASE)
        if not match:
            continue

        value = match.group(1).strip()
        if is_context_reference(value):
            return ""
        return value

    return ""


async def find_recent_text_message(channel) -> str | None:
    history = getattr(channel, "history", None)
    if history is None:
        return None

    async for message in history(limit=15):
        if message.author.bot:
            continue

        content = (message.content or "").strip()
        if not content:
            continue

        if content.startswith("/"):
            continue

        request_text = None
        if is_panam_addressed(content):
            request_text = re.sub(
                PANAM_STRIP_PREFIX_PATTERN,
                "",
                content,
                flags=re.IGNORECASE,
            ).strip()

        normalized = normalize_text(request_text or content)
        if not normalized or normalized == "panam":
            continue

        if (
            is_natural_attachment_analyze_request(normalized)
            or parse_natural_intent(normalized) is not None
        ):
            continue

        return content

    return None


def is_natural_analyze_request(content: str) -> bool:
    text = normalize_natural_text(content)
    return any(phrase in text for phrase in IMAGE_ANALYZE_TRIGGERS)


def is_natural_attachment_analyze_request(content: str) -> bool:
    if is_natural_analyze_request(content):
        return True

    text = normalize_natural_text(content)
    subject_pattern = "|".join(re.escape(subject) for subject in ATTACHMENT_SUBJECTS)

    return any(
        re.match(
            pattern_template.format(subject_pattern=subject_pattern),
            text,
            re.IGNORECASE,
        )
        for pattern_template in ATTACHMENT_ANALYZE_PATTERNS
    )


def is_generic_natural_attachment_request(content: str) -> bool:
    if is_natural_analyze_request(content):
        text = normalize_natural_text(content)
        return text in GENERIC_IMAGE_PHRASES

    text = normalize_natural_text(content)
    subject_pattern = "|".join(re.escape(subject) for subject in ATTACHMENT_SUBJECTS)
    return any(
        re.match(
            pattern_template.format(subject_pattern=subject_pattern),
            text,
            re.IGNORECASE,
        )
        for pattern_template in GENERIC_ATTACHMENT_PATTERNS
    )


def get_natural_analyze_question(content: str) -> str:
    text = normalize_natural_text(content)
    default_question = "Popiš, co je na obrázku."
    if text in GENERIC_IMAGE_PHRASES:
        return default_question

    return content.strip() or default_question


def get_natural_attachment_analyze_question(content: str) -> str:
    default_question = "Analyzuj tuto přílohu a stručně popiš, co obsahuje."
    if is_generic_natural_attachment_request(content):
        return default_question

    return content.strip() or default_question


def is_attachment_summary_request(text: str) -> bool:
    normalized = normalize_natural_text(text)
    return re.search(r"\b(?:shrn|precti)\b", normalized) is not None


def is_attachment_error_request(text: str) -> bool:
    normalized = normalize_natural_text(text)
    return any(
        phrase in normalized
        for phrase in (
            "chyba",
            "spatne",
            "problem",
            "najdi chybu",
            "co je tam spatne",
            "co je na tom spatne",
        )
    )


def is_attachment_content_request(text: str) -> bool:
    normalized = normalize_natural_text(text)
    return any(
        re.match(pattern, normalized) is not None
        for pattern in (
            r"^co\s+je\s+(?:v|ve|na)\s+(?:tom|toto|tohle).*$",
            r"^co\s+je\s+(?:v|ve)\s+(?:teto|te|ta)\s+priloze.*$",
            r"^co\s+tam\s+je.*$",
            r"^co\s+vidis.*$",
            r"^co\s+obsahuje(?:\s+(?:to|toto|tohle|tahle\s+priloha|ta\s+priloha))?.*$",
        )
    )


def is_attachment_look_request(text: str) -> bool:
    normalized = normalize_natural_text(text)
    return any(
        re.match(pattern, normalized) is not None
        for pattern in (
            r"^(?:koukni|mrkni)\s+na\s+(?:to|toto|tohle).*$",
            r"^podivej\s+se\s+na\s+(?:to|toto|tohle).*$",
            r"^analyzuj\s+(?:to|toto|tohle).*$",
        )
    )


def is_general_attachment_context_request(text: str) -> bool:
    return (
        is_attachment_summary_request(text)
        or is_attachment_error_request(text)
        or is_attachment_content_request(text)
        or is_attachment_look_request(text)
    )


def get_attachment_context_question(text: str, attachment_kind: str) -> str:
    if attachment_kind == "image" and is_attachment_error_request(text):
        return "Podívej se na obrázek a řekni, jaká chyba je tam vidět. Navrhni krátce další postup."

    if attachment_kind == "document" and is_attachment_summary_request(text):
        return "Shrň tuto přílohu."

    if attachment_kind == "document" and is_attachment_error_request(text):
        return "Najdi v dokumentu možné chyby nebo problémové části a stručně je vysvětli."

    return "Analyzuj tuto přílohu a stručně popiš, co obsahuje."


def is_context_attachment_request(text: str) -> bool:
    return is_general_attachment_context_request(text)


def get_context_attachment_question(text: str) -> str:
    return get_attachment_context_question(text, "document")


def is_file_router_candidate(text: str) -> bool:
    return has_explicit_file_output_request(text) or has_explicit_file_action_request(text)


def sanitize_classifier_context_text(text: str) -> str:
    normalized = normalize_natural_text(text)
    sensitive_markers = (
        "token",
        "api key",
        "apikey",
        "heslo",
        "password",
        "secret",
        "klic",
        "key",
    )
    if any(marker in normalized for marker in sensitive_markers):
        return "[sensitive content omitted]"

    return text[:500]


def get_recent_classifier_context(channel_id: int, limit: int = 4) -> list[dict]:
    recent_messages = panam_memory.get_messages(channel_id)[-limit:]
    context = []
    for message in recent_messages:
        role = message.get("role")
        content = str(message.get("content", "")).strip()
        if role not in {"user", "assistant"} or not content:
            continue

        context.append(
            {
                "role": role,
                "content": sanitize_classifier_context_text(content),
            }
        )

    return context


def get_natural_file_action_name(decision: dict) -> str:
    mode = decision.get("mode")
    if mode == "human_document":
        return "process_file"
    if mode == "structured_data":
        return "extract_data"
    if mode == "spreadsheet_transform":
        return "transform_excel"
    if mode == "unsupported_creative_spreadsheet_edit":
        return "unsupported_creative_spreadsheet_edit"
    if mode == "unsupported_direct_edit":
        return "unsupported_direct_edit"
    return "chat_answer"


async def send_source_message(source, text: str, file_path: Path | None = None) -> None:
    if isinstance(source, discord.Interaction):
        if file_path is None:
            await source.followup.send(text)
        else:
            await source.followup.send(text, file=discord.File(file_path))
        return

    if file_path is None:
        await source.reply(text, mention_author=False)
    else:
        await source.reply(text, file=discord.File(file_path), mention_author=False)


def mark_source_error(source) -> None:
    if isinstance(source, discord.Interaction):
        mark_command_status(source, "error")


async def run_human_document_file_job(
    source,
    file: discord.Attachment,
    instruction: str,
    output_format: str,
    action_type: str,
    action_name: str = "process_file",
) -> None:
    extension = get_file_extension(file.filename)
    normalized_output_format = output_format.lower().strip(".")
    output_extension = f".{normalized_output_format}"
    job = None

    try:
        context = get_safe_context(source)
        job = panam_files.create_file_job(
            user_id=get_context_int(context, "user_id"),
            channel_id=get_context_int(context, "channel_id"),
            action=action_name,
        )
        log_action(
            action_type,
            action_name,
            "started",
            source,
            job_id=job.job_id,
            filename=Path(file.filename).name,
            extension=extension,
            size_bytes=file.size,
            output_format=normalized_output_format,
        )

        input_path = await panam_files.save_attachment_to_job(file, job)
        log_action(
            action_type,
            action_name,
            "attachment_saved",
            source,
            job_id=job.job_id,
            filename=input_path.name,
            extension=input_path.suffix.lower(),
            size_bytes=input_path.stat().st_size,
            output_format=normalized_output_format,
        )

        data = input_path.read_bytes()
        extracted_text = extract_text_from_attachment(input_path.name, data)
        if not extracted_text.strip():
            mark_source_error(source)
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                action_type,
                action_name,
                "error",
                source,
                job_id=job.job_id,
                filename=input_path.name,
                extension=input_path.suffix.lower(),
                size_bytes=input_path.stat().st_size,
                output_format=normalized_output_format,
            )
            message_text = (
                "Z toho Excelu se mi nepodarilo vytahnout zadna data."
                if extension == ".xlsx"
                else "Z toho dokumentu se mi nepodarilo vytahnout zadny text."
            )
            await send_source_message(source, message_text)
            return

        extracted_text = trim_document_text(extracted_text)
        work_path = panam_files.write_work_text(job, "extracted_text.txt", extracted_text)
        log_action(
            action_type,
            action_name,
            "text_extracted",
            source,
            job_id=job.job_id,
            filename=work_path.name,
            extension=work_path.suffix.lower(),
            size_bytes=work_path.stat().st_size,
            output_format=normalized_output_format,
        )

        process_output_format = (
            "md" if normalized_output_format == "docx" else normalized_output_format
        )
        processed_text = await process_document_text(
            OPENAI_MODEL,
            extracted_text,
            instruction,
            input_path.name,
            process_output_format,
        )
        output_filename = panam_files.build_panam_output_filename(
            input_path.name,
            output_extension,
        )
        if normalized_output_format == "docx":
            output_path = panam_files.ensure_within_job(
                job.output_dir / output_filename,
                job,
            )
            panam_docx.create_docx_from_text(processed_text, output_path)
            job.output_files.append(
                {
                    "filename": output_path.name,
                    "extension": output_path.suffix.lower(),
                    "size_bytes": output_path.stat().st_size,
                }
            )
            panam_files.write_job_metadata(job)
        else:
            output_path = panam_files.write_output_text(
                job,
                output_filename,
                processed_text,
            )
        log_action(
            action_type,
            action_name,
            "output_written",
            source,
            job_id=job.job_id,
            filename=output_path.name,
            extension=output_path.suffix.lower(),
            size_bytes=output_path.stat().st_size,
            output_format=normalized_output_format,
        )

        job.status = "success"
        job.finished_at = datetime.now(timezone.utc).isoformat()
        panam_files.write_job_metadata(job)

        await send_source_message(source, "Soubor je zpracovany.", output_path)
        set_last_file_context(
            context.get("channel_id"),
            Path(file.filename).name,
            extension,
            "human_document",
            output_path.name,
        )
        log_action(
            action_type,
            action_name,
            "success",
            source,
            job_id=job.job_id,
            filename=Path(file.filename).name,
            extension=extension,
            size_bytes=file.size,
            output_format=normalized_output_format,
        )

    except AttachmentAnalysisUserError as error:
        mark_source_error(source)
        if job is not None:
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                action_type,
                action_name,
                "error",
                source,
                job_id=job.job_id,
                filename=Path(file.filename).name,
                extension=extension,
                size_bytes=file.size,
                output_format=normalized_output_format,
            )
        logger.exception("Chyba pri zpracovani file-job pipeline")
        await send_source_message(source, str(error))

    except Exception:
        mark_source_error(source)
        if job is not None:
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                action_type,
                action_name,
                "error",
                source,
                job_id=job.job_id,
                filename=Path(file.filename).name,
                extension=extension,
                size_bytes=file.size,
                output_format=normalized_output_format,
            )
        logger.exception("Chyba pri zpracovani file-job pipeline")
        await send_source_message(source, "Neco se pokazilo pri zpracovani souboru.")

    finally:
        if job is not None:
            panam_files.cleanup_job(job)


async def run_structured_data_file_job(
    source,
    file: discord.Attachment,
    instruction: str,
    output_format: str,
    action_type: str,
    action_name: str = "extract_data",
) -> None:
    extension = get_file_extension(file.filename)
    normalized_output_format = output_format.lower().strip(".")
    output_extension = f".{normalized_output_format}"
    job = None

    try:
        context = get_safe_context(source)
        job = panam_files.create_file_job(
            user_id=get_context_int(context, "user_id"),
            channel_id=get_context_int(context, "channel_id"),
            action=action_name,
        )
        log_action(
            action_type,
            action_name,
            "started",
            source,
            job_id=job.job_id,
            filename=Path(file.filename).name,
            extension=extension,
            size_bytes=file.size,
            output_format=normalized_output_format,
        )

        input_path = await panam_files.save_attachment_to_job(file, job)
        log_action(
            action_type,
            action_name,
            "attachment_saved",
            source,
            job_id=job.job_id,
            filename=input_path.name,
            extension=input_path.suffix.lower(),
            size_bytes=input_path.stat().st_size,
            output_format=normalized_output_format,
        )

        data = input_path.read_bytes()
        extracted_text = extract_text_from_attachment(input_path.name, data)
        if not extracted_text.strip():
            mark_source_error(source)
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                action_type,
                action_name,
                "error",
                source,
                job_id=job.job_id,
                filename=input_path.name,
                extension=input_path.suffix.lower(),
                size_bytes=input_path.stat().st_size,
                output_format=normalized_output_format,
            )
            message_text = (
                "Z toho Excelu se mi nepodarilo vytahnout zadna data."
                if extension == ".xlsx"
                else "Z toho dokumentu se mi nepodarilo vytahnout zadny text."
            )
            await send_source_message(source, message_text)
            return

        extracted_text = trim_document_text(extracted_text)
        work_path = panam_files.write_work_text(job, "extracted_text.txt", extracted_text)
        log_action(
            action_type,
            action_name,
            "text_extracted",
            source,
            job_id=job.job_id,
            filename=work_path.name,
            extension=work_path.suffix.lower(),
            size_bytes=work_path.stat().st_size,
            output_format=normalized_output_format,
        )

        structured_text = await extract_structured_data(
            OPENAI_MODEL,
            extracted_text,
            instruction,
            input_path.name,
            "json" if normalized_output_format == "xlsx" else normalized_output_format,
        )
        log_action(
            action_type,
            action_name,
            "ai_extraction_completed",
            source,
            job_id=job.job_id,
            filename=input_path.name,
            extension=input_path.suffix.lower(),
            size_bytes=input_path.stat().st_size,
            output_format=normalized_output_format,
        )

        json_data = None
        if normalized_output_format in {"json", "xlsx"}:
            try:
                json_data = json.loads(structured_text)
            except json.JSONDecodeError:
                mark_source_error(source)
                job.status = "error"
                job.finished_at = datetime.now(timezone.utc).isoformat()
                panam_files.write_job_metadata(job)
                log_action(
                    action_type,
                    action_name,
                    "error",
                    source,
                    job_id=job.job_id,
                    filename=input_path.name,
                    extension=input_path.suffix.lower(),
                    size_bytes=input_path.stat().st_size,
                    output_format=normalized_output_format,
                )
                logger.warning(
                    "extract_data JSON validation failed job_id=%s status=error",
                    job.job_id,
                )
                message_text = (
                    "Data pro XLSX se nepodarilo validovat. Zkus presnejsi instrukci nebo output_format=json."
                    if normalized_output_format == "xlsx"
                    else "Vystup JSON se nepodarilo validovat. Zkus presnejsi instrukci nebo output_format=md."
                )
                await send_source_message(source, message_text)
                return
        elif normalized_output_format == "csv":
            if not structured_text.strip() or not structured_text.strip().splitlines():
                raise ValueError("CSV output is empty.")
        elif not structured_text.strip():
            raise ValueError("Markdown output is empty.")

        output_filename = panam_files.build_panam_output_filename(
            input_path.name,
            output_extension,
        )

        if normalized_output_format == "xlsx":
            output_path = panam_files.ensure_within_job(
                job.output_dir / output_filename,
                job,
            )
            panam_excel.create_xlsx_from_items(
                get_items_for_xlsx(json_data),
                output_path,
            )
            job.output_files.append(
                {
                    "filename": output_path.name,
                    "extension": output_path.suffix.lower(),
                    "size_bytes": output_path.stat().st_size,
                }
            )
            panam_files.write_job_metadata(job)
            output_status = "xlsx_output_written"
        else:
            output_path = panam_files.write_output_text(
                job,
                output_filename,
                structured_text,
            )
            output_status = "output_written"

        log_action(
            action_type,
            action_name,
            output_status,
            source,
            job_id=job.job_id,
            filename=output_path.name,
            extension=output_path.suffix.lower(),
            size_bytes=output_path.stat().st_size,
            output_format=normalized_output_format,
        )

        job.status = "success"
        job.finished_at = datetime.now(timezone.utc).isoformat()
        panam_files.write_job_metadata(job)

        await send_source_message(source, "Data jsou vytezena.", output_path)
        set_last_file_context(
            context.get("channel_id"),
            Path(file.filename).name,
            extension,
            "structured_data",
            output_path.name,
        )
        log_action(
            action_type,
            action_name,
            "success",
            source,
            job_id=job.job_id,
            filename=Path(file.filename).name,
            extension=extension,
            size_bytes=file.size,
            output_format=normalized_output_format,
        )

    except AttachmentAnalysisUserError as error:
        mark_source_error(source)
        if job is not None:
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                action_type,
                action_name,
                "error",
                source,
                job_id=job.job_id,
                filename=Path(file.filename).name,
                extension=extension,
                size_bytes=file.size,
                output_format=normalized_output_format,
            )
        logger.exception("Chyba pri tezeni dat ze souboru")
        await send_source_message(source, str(error))

    except Exception:
        mark_source_error(source)
        if job is not None:
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                action_type,
                action_name,
                "error",
                source,
                job_id=job.job_id,
                filename=Path(file.filename).name,
                extension=extension,
                size_bytes=file.size,
                output_format=normalized_output_format,
            )
        logger.exception("Chyba pri tezeni dat ze souboru")
        await send_source_message(source, "Neco se pokazilo pri tezeni dat ze souboru.")

    finally:
        if job is not None:
            panam_files.cleanup_job(job)


async def run_spreadsheet_transform_file_job(
    source,
    file: discord.Attachment,
    instruction: str,
    action_type: str,
    action_name: str = "transform_excel",
) -> None:
    extension = get_file_extension(file.filename)
    job = None

    try:
        context = get_safe_context(source)
        job = panam_files.create_file_job(
            user_id=get_context_int(context, "user_id"),
            channel_id=get_context_int(context, "channel_id"),
            action="spreadsheet_transform",
        )
        log_action(
            action_type,
            action_name,
            "started",
            source,
            job_id=job.job_id,
            filename=Path(file.filename).name,
            extension=extension,
            size_bytes=file.size,
            output_format="xlsx",
        )

        input_path = await panam_files.save_attachment_to_job(file, job)
        log_action(
            action_type,
            action_name,
            "attachment_saved",
            source,
            job_id=job.job_id,
            filename=input_path.name,
            extension=input_path.suffix.lower(),
            size_bytes=input_path.stat().st_size,
            output_format="xlsx",
        )

        output_filename = panam_files.build_panam_output_filename(
            input_path.name,
            ".xlsx",
        )
        output_path = panam_files.ensure_within_job(
            job.output_dir / output_filename,
            job,
        )
        result = panam_spreadsheet.transform_xlsx(
            input_path,
            output_path,
            instruction,
        )
        job.output_files.append(
            {
                "filename": output_path.name,
                "extension": output_path.suffix.lower(),
                "size_bytes": output_path.stat().st_size,
            }
        )
        panam_files.write_job_metadata(job)
        log_action(
            action_type,
            action_name,
            "xlsx_output_written",
            source,
            job_id=job.job_id,
            filename=output_path.name,
            extension=output_path.suffix.lower(),
            size_bytes=output_path.stat().st_size,
            output_format="xlsx",
            rows_read=result.rows_read,
            rows_written=result.rows_written,
            columns_written=result.columns_written,
            operations=",".join(result.operations),
        )

        job.status = "success"
        job.finished_at = datetime.now(timezone.utc).isoformat()
        panam_files.write_job_metadata(job)

        success_message_parts = [
            "Excel jsem upravila jako nový soubor.",
            f"Provedená změna: {result.operation_summary}.",
        ]
        if (
            result.row_count_before is not None
            and result.row_count_after is not None
        ):
            success_message_parts.append(
                f"Řádky: {result.row_count_before} -> {result.row_count_after}."
            )
        if (
            result.column_count_before is not None
            and result.column_count_after is not None
            and result.column_count_before != result.column_count_after
        ):
            success_message_parts.append(
                f"Sloupce: {result.column_count_before} -> {result.column_count_after}."
            )
        success_message_parts.append("Původní příloha zůstala beze změn.")
        await send_source_message(
            source,
            " ".join(success_message_parts),
            result.output_path,
        )
        set_last_file_context(
            context.get("channel_id"),
            Path(file.filename).name,
            extension,
            "spreadsheet_transform",
            output_path.name,
        )
        log_action(
            action_type,
            action_name,
            "success",
            source,
            job_id=job.job_id,
            filename=Path(file.filename).name,
            extension=extension,
            size_bytes=file.size,
            output_format="xlsx",
            rows_read=result.rows_read,
            rows_written=result.rows_written,
            columns_written=result.columns_written,
            operations=",".join(result.operations),
        )

    except panam_spreadsheet.SpreadsheetTransformUserError as error:
        mark_source_error(source)
        if job is not None:
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                action_type,
                action_name,
                "error",
                source,
                job_id=job.job_id,
                filename=Path(file.filename).name,
                extension=extension,
                size_bytes=file.size,
                output_format="xlsx",
            )
        logger.warning(
            "spreadsheet_transform user_error job_id=%s",
            job.job_id if job else None,
        )
        await send_source_message(source, str(error))

    except Exception:
        mark_source_error(source)
        if job is not None:
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                action_type,
                action_name,
                "error",
                source,
                job_id=job.job_id,
                filename=Path(file.filename).name,
                extension=extension,
                size_bytes=file.size,
                output_format="xlsx",
            )
        logger.exception("Chyba pri spreadsheet_transform file-job pipeline")
        await send_source_message(source, "Neco se pokazilo pri uprave Excelu.")

    finally:
        if job is not None:
            panam_files.cleanup_job(job)


def remember_router_decision(
    message: discord.Message,
    decision: dict,
    default_target: str = "current_attachment",
) -> None:
    mode = str(decision.get("mode") or "chat_answer")
    target = str(decision.get("target") or default_target)
    output_format = decision.get("output_format")
    if not isinstance(output_format, str):
        output_format = None

    confidence = decision.get("confidence")
    if not isinstance(confidence, (int, float)):
        confidence = None

    set_last_router_decision(
        message.channel.id,
        target,
        mode,
        output_format=output_format,
        classifier_used=bool(decision.get("classifier_used", False)),
        confidence=confidence,
    )


async def handle_natural_file_request(
    message: discord.Message,
    file: discord.Attachment | None,
    decision: dict,
) -> None:
    mode = decision.get("mode")
    action_name = get_natural_file_action_name(decision)

    if mode == "unsupported_creative_spreadsheet_edit":
        remember_router_decision(
            message,
            decision,
            "current_attachment" if file is not None else "none",
        )
        log_action("natural_message", action_name, "started", message)
        await message.reply(
            CREATIVE_SPREADSHEET_FALLBACK_MESSAGE,
            mention_author=False,
        )
        log_action("natural_message", action_name, "success", message)
        return

    if mode == "unsupported_direct_edit":
        remember_router_decision(
            message,
            decision,
            "current_attachment" if file is not None else "none",
        )
        log_action("natural_message", action_name, "started", message)
        await message.reply(
            "Puvodni Excel ani puvodni prilohu zatim primo neupravuju. "
            "Muzu ale vytvorit novy XLSX vystup, kdyz pozadas o podporovanou jednoduchou upravu.\n\n"
            f"{panam_spreadsheet.unsupported_instruction_message()}",
            mention_author=False,
        )
        log_action("natural_message", action_name, "success", message)
        return

    if file is None:
        remember_router_decision(message, decision, "none")
        await message.reply(
            "Nevidim zadnou podporovanou prilohu ani v teto zprave, ani v predchozich zpravach.",
            mention_author=False,
        )
        return

    attachment_details = get_safe_attachment_info(file)
    attachment_details["kind"] = attachment_details.pop("file_type", None)

    if mode == "chat_answer":
        log_action(
            "natural_message",
            action_name,
            "started",
            message,
            **attachment_details,
        )
        try:
            question = decision.get("question") or get_attachment_context_question(
                message.content or "",
                get_attachment_kind(file) or "document",
            )
            answer = await analyze_selected_attachment(file, question)
            await message.reply(answer, mention_author=False)
            set_last_file_context(
                message.channel.id,
                Path(file.filename).name,
                get_file_extension(file.filename),
                "chat_answer",
                file_summary=answer,
            )
            remember_router_decision(message, decision)
            log_action(
                "natural_message",
                action_name,
                "success",
                message,
                **attachment_details,
            )
            return

        except AttachmentAnalysisUserError as error:
            await message.reply(str(error), mention_author=False)
            log_action(
                "natural_message",
                action_name,
                "error",
                message,
                **attachment_details,
            )
            return

        except Exception:
            log_action(
                "natural_message",
                action_name,
                "error",
                message,
                **attachment_details,
            )
            logger.exception("Chyba pri zpracovani natural file chat odpovedi")
            await message.reply(
                "Neco se pokazilo pri analyze prilohy. Mrkni do konzole na chybu.",
                mention_author=False,
            )
            return

    if get_attachment_kind(file) != "document":
        await message.reply(
            "Vystupni soubory zatim umim tvorit jen z dokumentu TXT, MD, CSV, PDF, DOCX nebo XLSX.",
            mention_author=False,
        )
        return

    extension = get_file_extension(file.filename)
    if file.size > MAX_DOCUMENT_SIZE_BYTES:
        await message.reply("Ten soubor je moc velky. Zatim beru max 20 MB.", mention_author=False)
        return

    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        await message.reply(
            "Tenhle rezim podporuje dokumenty TXT, MD, CSV, PDF, DOCX nebo XLSX.",
            mention_author=False,
        )
        return

    async with message.channel.typing():
        if mode == "spreadsheet_transform":
            if extension != ".xlsx":
                remember_router_decision(message, decision)
                await message.reply(
                    "Excel transform v1 podporuje jen XLSX soubory. Puvodni prilohu neupravuju.",
                    mention_author=False,
                )
                return

            decision["output_format"] = "xlsx"
            remember_router_decision(message, decision)
            await run_spreadsheet_transform_file_job(
                message,
                file,
                decision.get("instruction", message.content or ""),
                "natural_message",
            )
            return

        if mode == "human_document":
            output_format = decision.get("output_format", "md")
            if output_format not in {"md", "txt", "docx"}:
                output_format = "md"

            decision["output_format"] = output_format
            remember_router_decision(message, decision)
            await run_human_document_file_job(
                message,
                file,
                decision.get("instruction", message.content or ""),
                output_format,
                "natural_message",
            )
            return

        if mode == "structured_data":
            output_format = decision.get("output_format", "json")
            if output_format not in {"json", "csv", "md", "xlsx"}:
                output_format = "json"

            decision["output_format"] = output_format
            remember_router_decision(message, decision)
            await run_structured_data_file_job(
                message,
                file,
                decision.get("instruction", message.content or ""),
                output_format,
                "natural_message",
            )
            return

    await message.reply(
        "Tohle jsem u prilohy nepoznala, tak volim bezpecnou odpoved do chatu.",
        mention_author=False,
    )


async def handle_ai_classified_file_request(
    message: discord.Message,
    request_text: str,
    current_file: discord.Attachment | None,
) -> bool:
    last_file_context = get_last_file_context(message.channel.id)
    recent_context = get_recent_classifier_context(message.channel.id)

    if not (current_file is not None or last_file_context is not None or recent_context):
        log_action(
            "natural_message",
            "ai_file_intent_classifier",
            "skipped",
            message,
            classifier_used=False,
            target=None,
            mode=None,
            output_format=None,
            confidence=None,
            classifier_status="no_context",
        )
        return False

    current_attachment_context = {}
    if current_file is not None:
        current_attachment_context = {
            "exists": True,
            "filename": Path(current_file.filename).name,
            "extension": get_file_extension(current_file.filename),
        }
    else:
        current_attachment_context = {"exists": False}

    file_context = {
        "current_attachment": current_attachment_context,
        "last_file_context": last_file_context,
    }

    try:
        raw_intent = await classify_file_request_intent(
            OPENAI_MODEL,
            request_text,
            file_context=file_context,
            recent_context=recent_context,
        )
        intent = validate_ai_file_intent(
            raw_intent,
            has_current_attachment=current_file is not None,
            has_last_file_context=last_file_context is not None,
        )
    except Exception:
        log_action(
            "natural_message",
            "ai_file_intent_classifier",
            "error",
            message,
            classifier_used=True,
            target=None,
            mode=None,
            output_format=None,
            confidence=None,
            classifier_status="error",
        )
        logger.exception("Chyba pri AI klasifikaci nejasneho file/context dotazu")
        return False

    classifier_status = (
        "fallback"
        if intent.get("target") in {"conversation", "none"} or float(intent.get("confidence", 0.0)) < 0.65
        else "selected"
    )
    log_action(
        "natural_message",
        "ai_file_intent_classifier",
        "success",
        message,
        classifier_used=True,
        target=intent.get("target"),
        mode=intent.get("mode"),
        output_format=intent.get("output_format"),
        confidence=round(float(intent.get("confidence", 0.0)), 2),
        classifier_status=classifier_status,
    )
    set_last_router_decision(
        message.channel.id,
        str(intent.get("target") or "none"),
        str(intent.get("mode") or "chat_answer"),
        output_format=intent.get("output_format") if isinstance(intent.get("output_format"), str) else None,
        classifier_used=True,
        confidence=float(intent.get("confidence", 0.0)),
    )

    target = intent.get("target")
    mode = intent.get("mode")
    if target in {"conversation", "none"}:
        return False

    if mode == "unsupported_direct_edit":
        await message.reply(
            "Puvodni prilohu zatim primo neupravuju. "
            "Muzu ale vytvorit novy XLSX vystup, kdyz pozadas o podporovanou jednoduchou upravu.\n\n"
            f"{panam_spreadsheet.unsupported_instruction_message()}",
            mention_author=False,
        )
        return True

    if mode == "unsupported_creative_spreadsheet_edit":
        await message.reply(
            CREATIVE_SPREADSHEET_FALLBACK_MESSAGE,
            mention_author=False,
        )
        return True

    selected_file = current_file
    if target == "last_file_context":
        selected_file = await find_recent_supported_attachment(message.channel)

    if selected_file is None:
        return False

    if mode == "human_document":
        decision = {
            "target": target,
            "mode": "human_document",
            "instruction": intent.get("instruction") or request_text,
            "output_format": intent.get("output_format") or "md",
            "classifier_used": True,
            "confidence": intent.get("confidence"),
        }
    elif mode == "structured_data":
        decision = {
            "target": target,
            "mode": "structured_data",
            "instruction": intent.get("instruction") or request_text,
            "output_format": intent.get("output_format") or "json",
            "classifier_used": True,
            "confidence": intent.get("confidence"),
        }
    elif mode == "spreadsheet_transform":
        decision = {
            "target": target,
            "mode": "spreadsheet_transform",
            "instruction": intent.get("instruction") or request_text,
            "output_format": "xlsx",
            "classifier_used": True,
            "confidence": intent.get("confidence"),
        }
    else:
        decision = {
            "target": target,
            "mode": "chat_answer",
            "question": intent.get("question") or get_attachment_context_question(
                request_text,
                get_attachment_kind(selected_file) or "document",
            ),
            "classifier_used": True,
            "confidence": intent.get("confidence"),
        }

    await handle_natural_file_request(message, selected_file, decision)
    return True


async def handle_file_summary_followup(
    message: discord.Message,
    request_text: str,
    file_context: dict,
) -> bool:
    file_summary = file_context.get("file_summary")
    if not isinstance(file_summary, str) or not file_summary.strip():
        return False

    log_action("natural_message", "file_summary_followup", "started", message)
    try:
        prompt = (
            "Uzivatel navazuje na posledni zpracovany soubor. "
            "Bezpecne shrnuti souboru: "
            f"{file_summary.strip()}\n\n"
            "Odpovez podle tohoto shrnuti. Nepredstirej, ze znas cely obsah souboru. "
            "Pokud z tohoto shrnuti nejde odpovedet, rekni to kratce a pozadej o soubor nebo upresneni.\n\n"
            f"Dotaz uzivatele: {request_text.strip()}"
        )
        answer = await ask_panam(OPENAI_MODEL, prompt)
        await message.reply(answer, mention_author=False)
        panam_memory.add_message(message.channel.id, "user", message.content or request_text)
        panam_memory.add_message(message.channel.id, "assistant", answer)
        log_action("natural_message", "file_summary_followup", "success", message)
        return True

    except Exception:
        log_action("natural_message", "file_summary_followup", "error", message)
        logger.exception("Chyba pri odpovedi podle bezpecneho file summary")
        await message.reply(
            "Mam ulozene jen kratke shrnuti posledniho souboru, ale ted se mi z nej nepodarilo odpovedet.",
            mention_author=False,
        )
        return True


def get_items_for_xlsx(data) -> list[dict]:
    if isinstance(data, dict) and isinstance(data.get("items"), list):
        raw_items = data["items"]
    elif isinstance(data, list):
        raw_items = data
    elif isinstance(data, dict):
        raw_items = [data]
    else:
        raw_items = [{"value": data}]

    items = []
    for item in raw_items:
        if isinstance(item, dict):
            items.append(item)
        else:
            items.append({"value": item})

    return items


def parse_natural_intent(text: str) -> Optional[tuple[str, Optional[str]]]:
    for pattern in HELP_PATTERNS:
        if re.match(pattern, text, re.IGNORECASE):
            return "help", None

    if re.match(NOTE_ADD_PREVIOUS_PATTERN, text, re.IGNORECASE):
        return "note_add_previous", None

    for pattern in OPINION_CONTEXT_PATTERNS:
        if re.match(pattern, text, re.IGNORECASE):
            return "ask_previous", None

    if re.match(SUMMARY_CONTEXT_PATTERN, text, re.IGNORECASE):
        return "summary_previous", None

    for intent, pattern in NATURAL_INTENT_PATTERNS:
        match = re.match(pattern, text, re.IGNORECASE)
        if match:
            return intent, match.group(1).strip()

    if re.match(NOTE_LIST_PATTERN, text, re.IGNORECASE):
        return "note_list", None

    if re.match(TODO_LIST_PATTERN, text, re.IGNORECASE):
        return "todo_list", None

    return None


def get_natural_action_name(request_text: str, original_content: str) -> str | None:
    file_decision = decide_file_response_mode(request_text)
    if is_file_router_candidate(request_text):
        return get_natural_file_action_name(file_decision)

    if is_general_attachment_context_request(request_text) and has_explicit_file_action_request(request_text):
        return "natural_attachment_context"

    if is_natural_attachment_analyze_request(request_text) and has_explicit_file_action_request(request_text):
        return "analyze_attachment"

    intent = parse_natural_intent(request_text)
    if intent is not None:
        intent_name, _ = intent
        if intent_name in ("ask_previous", "ask"):
            return "ask"
        if intent_name in ("summary_previous", "summary"):
            return "summary"
        if intent_name in ("note_add_previous", "note_add"):
            return "note_add"
        return intent_name

    if extract_basic_panam_prompt(original_content) is not None:
        return "ask"

    return None


class DiscordAIBot(discord.Client):
    def __init__(self) -> None:
        intents = discord.Intents.default()
        intents.message_content = True

        # Message content intent je potřeba pro práci s obsahem zpráv.
        super().__init__(intents=intents)

        self.tree = app_commands.CommandTree(self)

    async def setup_hook(self) -> None:
        """
        Registrace slash commandů.

        Pro PoC doporučuji použít DISCORD_GUILD_IDS.
        Guild sync je rychlý, často prakticky hned.
        Globální sync může trvat déle.
        """
        if DISCORD_GUILD_IDS:
            for guild_id in DISCORD_GUILD_IDS:
                guild = discord.Object(id=int(guild_id))
                self.tree.copy_global_to(guild=guild)
                await self.tree.sync(guild=guild)

            logger.info(
                "Slash commandy synchronizovány pro %s serverů.",
                len(DISCORD_GUILD_IDS),
            )
        else:
            await self.tree.sync()
            logger.info("Slash commandy synchronizovány globálně.")

    async def on_ready(self) -> None:
        if self.user is None:
            logger.info("Bot je online, ale user zatím není dostupný.")
            return

        logger.info("Bot je online jako %s | ID: %s", self.user, self.user.id)

    async def on_message(self, message: discord.Message) -> None:
        if message.author.bot:
            return

        if self.user is None:
            return

        request_text = extract_panam_request(message, self.user)
        if request_text is None:
            return

        current_file = find_supported_attachment_in_message(message)
        selected_file = current_file
        natural_action = get_natural_action_name(request_text, message.content or "")
        if ALLOWED_CHANNEL_IDS and str(message.channel.id) not in ALLOWED_CHANNEL_IDS:
            if natural_action is not None:
                log_action("natural_message", natural_action, "denied", message)
            return

        explicit_file_output_request = has_explicit_file_output_request(request_text)
        explicit_file_action_request = has_explicit_file_action_request(request_text)
        meta_router_discussion = is_meta_router_or_behavior_discussion(request_text)
        direct_file_edit_request = has_direct_file_edit_request(request_text)
        last_file_context = get_last_file_context(message.channel.id)
        last_context_extension = (
            str(last_file_context.get("source_extension") or "").lower()
            if last_file_context is not None
            else ""
        )
        has_xlsx_file_context = (
            get_file_extension(current_file.filename) == ".xlsx"
            if current_file is not None
            else last_context_extension == ".xlsx"
        )
        spreadsheet_transform_candidate = has_spreadsheet_transform_request(
            request_text,
            get_file_extension(current_file.filename) if current_file is not None else last_context_extension,
            has_xlsx_context=has_xlsx_file_context,
        )
        unsafe_creative_spreadsheet_candidate = has_unsafe_creative_spreadsheet_request(
            request_text,
            get_file_extension(current_file.filename) if current_file is not None else last_context_extension,
            has_xlsx_context=has_xlsx_file_context,
        )
        last_file_reference = (
            current_file is None
            and not meta_router_discussion
            and matches_last_file_reference(request_text, last_file_context)
        )
        file_router_candidate = (
            explicit_file_output_request
            or explicit_file_action_request
            or spreadsheet_transform_candidate
            or unsafe_creative_spreadsheet_candidate
            or direct_file_edit_request
            or last_file_reference
        )
        if file_router_candidate and selected_file is None:
            selected_file = (
                await find_recent_xlsx_attachment(message.channel)
                if spreadsheet_transform_candidate or unsafe_creative_spreadsheet_candidate
                else await find_recent_supported_attachment(message.channel)
            )

        file_decision = decide_file_response_mode(
            request_text,
            get_file_extension(selected_file.filename) if selected_file is not None else None,
            has_xlsx_context=has_xlsx_file_context,
        )

        if (
            current_file is None
            and not meta_router_discussion
            and last_file_context is not None
            and last_file_context.get("file_summary")
            and last_file_reference
            and is_followup_to_file_summary(request_text)
            and file_decision.get("mode") == "chat_answer"
        ):
            handled_by_summary = await handle_file_summary_followup(
                message,
                request_text,
                last_file_context,
            )
            if handled_by_summary:
                return

        if file_router_candidate and (
            selected_file is not None
            or file_decision.get("mode") in {"human_document", "structured_data", "spreadsheet_transform", "unsupported_creative_spreadsheet_edit", "unsupported_direct_edit"}
            or explicit_file_output_request
            or explicit_file_action_request
            or spreadsheet_transform_candidate
            or unsafe_creative_spreadsheet_candidate
            or last_file_reference
        ):
            file_decision["target"] = (
                "current_attachment"
                if current_file is not None
                else "last_file_context"
                if selected_file is not None
                else "none"
            )
            await handle_natural_file_request(message, selected_file, file_decision)
            return

        if (
            current_file is None
            and not meta_router_discussion
            and is_followup_to_file_summary(request_text)
        ):
            if last_file_context is not None and last_file_context.get("file_summary"):
                handled_by_summary = await handle_file_summary_followup(
                    message,
                    request_text,
                    last_file_context,
                )
                if handled_by_summary:
                    return

        if is_ambiguous_context_request(request_text) and not meta_router_discussion:
            handled_by_classifier = await handle_ai_classified_file_request(
                message,
                request_text,
                current_file,
            )
            if handled_by_classifier:
                return

        context_attachment_request = (
            selected_file is not None
            and not is_ambiguous_context_request(request_text)
            and not meta_router_discussion
            and is_general_attachment_context_request(request_text)
        )
        if context_attachment_request and selected_file is None:
            selected_file = await find_recent_supported_attachment(message.channel)

        if context_attachment_request and selected_file is not None:
            attachment_details = {}
            if selected_file is not None:
                attachment_details = get_safe_attachment_info(selected_file)
                attachment_details["kind"] = attachment_details.pop("file_type", None)

            log_action(
                "natural_message",
                "natural_attachment_context",
                "started",
                message,
                **attachment_details,
            )
            try:
                attachment_kind = get_attachment_kind(selected_file) or "document"
                question = get_attachment_context_question(request_text, attachment_kind)
                answer = await analyze_selected_attachment(selected_file, question)
                await message.reply(answer, mention_author=False)
                set_last_file_context(
                    message.channel.id,
                    Path(selected_file.filename).name,
                    get_file_extension(selected_file.filename),
                    "chat_answer",
                    file_summary=answer,
                )
                remember_router_decision(
                    message,
                    {
                        "target": "current_attachment" if current_file is not None else "last_file_context",
                        "mode": "chat_answer",
                    },
                )
                log_action(
                    "natural_message",
                    "natural_attachment_context",
                    "success",
                    message,
                    **attachment_details,
                )
                return

            except AttachmentAnalysisUserError as error:
                await message.reply(
                    str(error),
                    mention_author=False,
                )
                log_action(
                    "natural_message",
                    "natural_attachment_context",
                    "error",
                    message,
                    **attachment_details,
                )
                return

            except Exception:
                log_action(
                    "natural_message",
                    "natural_attachment_context",
                    "error",
                    message,
                    **attachment_details,
                )
                logger.exception("Chyba při zpracování přirozené kontextové přílohy")
                await message.reply(
                    "Něco se pokazilo při analýze přílohy. Mrkni do konzole na chybu.",
                    mention_author=False,
                )
                return

        if context_attachment_request and selected_file is None:
            fallback_intent = parse_natural_intent(request_text)
            if fallback_intent is None:
                basic_prompt = extract_basic_panam_prompt(message.content or "")
                natural_action = "ask" if basic_prompt is not None else None
            else:
                fallback_intent_name, _ = fallback_intent
                if fallback_intent_name in ("ask_previous", "ask"):
                    natural_action = "ask"
                elif fallback_intent_name in ("summary_previous", "summary"):
                    natural_action = "summary"
                elif fallback_intent_name in ("note_add_previous", "note_add"):
                    natural_action = "note_add"
                else:
                    natural_action = fallback_intent_name

        if is_natural_attachment_analyze_request(request_text) and (
            selected_file is not None or explicit_file_action_request
        ):
            log_action("natural_message", "analyze_attachment", "started", message)
            try:
                if selected_file is None:
                    selected_file = await find_recent_supported_attachment(message.channel)

                if selected_file is None:
                    await message.reply(
                        "Nevidím žádnou podporovanou přílohu ani v této zprávě, ani v předchozích zprávách.",
                        mention_author=False,
                    )
                    log_action("natural_message", "analyze_attachment", "success", message)
                    return

                question = get_natural_attachment_analyze_question(request_text)
                answer = await analyze_selected_attachment(selected_file, question)
                await message.reply(answer, mention_author=False)
                log_action(
                    "natural_message",
                    "analyze_attachment",
                    "success",
                    message,
                    **get_safe_attachment_info(selected_file),
                )
                return

            except AttachmentAnalysisUserError as error:
                await message.reply(
                    str(error),
                    mention_author=False,
                )
                log_action("natural_message", "analyze_attachment", "error", message)
                return

            except Exception:
                log_action("natural_message", "analyze_attachment", "error", message)
                logger.exception("Chyba při zpracování přirozené analýzy přílohy")
                await message.reply(
                    "Něco se pokazilo při analýze přílohy. Mrkni do konzole na chybu.",
                    mention_author=False,
                )
                return

        intent = parse_natural_intent(request_text)
        if intent is None:
            basic_prompt = extract_basic_panam_prompt(message.content or "")
            if basic_prompt is not None:
                log_action("natural_message", "ask", "started", message)
                try:
                    await handle_basic_panam_message(message, basic_prompt)
                except Exception:
                    log_action("natural_message", "ask", "error", message)
                    logger.exception("Chyba při zpracování přirozené zprávy")
                    await message.channel.send("Něco se pokazilo při zpracování akce.")
                else:
                    log_action("natural_message", "ask", "success", message)
                return

            await message.channel.send("Tohle zatím neumím převést na akci. Zkus /help.")
            return

        intent_name, value = intent
        natural_action = natural_action or intent_name
        natural_status = "success"
        log_action("natural_message", natural_action, "started", message)

        try:
            if intent_name == "help":
                await message.channel.send(get_help_text())
                return

            if intent_name == "note_add_previous":
                previous_content = await find_recent_text_message(message.channel)
                if previous_content is None:
                    await message.channel.send(
                        "Nemám co uložit. Napiš text poznámky nebo to pošli pod zprávu, kterou si mám zapamatovat."
                    )
                    return

                create_note(previous_content, message.author, message.channel.id)
                await message.channel.send("Poznámka uložená.")
                return

            if intent_name == "note_add" and value:
                inline_note = extract_inline_content_after_trigger(
                    request_text,
                    list(NOTE_ADD_TRIGGERS),
                )
                if inline_note:
                    value = inline_note

                if is_context_reference(value):
                    previous_content = await find_recent_text_message(message.channel)
                    if previous_content is None:
                        await message.channel.send(
                            "Nemám co uložit. Napiš text poznámky nebo to pošli pod zprávu, kterou si mám zapamatovat."
                        )
                        return

                    create_note(previous_content, message.author, message.channel.id)
                    await message.channel.send("Poznámka uložená.")
                    return

                create_note(value, message.author, message.channel.id)
                await message.channel.send("Poznámka uložená.")
                return

            if intent_name == "todo_add" and value:
                todo_id = create_todo(value, message.author, message.channel.id)
                await message.channel.send(f"Úkol #{todo_id} uložený.")
                return

            if intent_name == "note_search" and value:
                await message.channel.send(format_note_search_response(value))
                return

            if intent_name == "note_list":
                await message.channel.send(format_note_list_response())
                return

            if intent_name == "todo_list":
                await message.channel.send(format_todo_list_response())
                return

            if intent_name == "ask" and value:
                inline_question = extract_inline_content_after_trigger(
                    request_text,
                    list(OPINION_TRIGGERS),
                )
                if inline_question:
                    value = inline_question

                if is_context_reference(value):
                    selected_file = find_supported_attachment_in_message(message)
                    if selected_file is not None:
                        answer = await analyze_selected_attachment(
                            selected_file,
                            "Co si o této příloze myslíš?",
                        )
                        await message.reply(answer, mention_author=False)
                        return

                    previous_content = await find_recent_text_message(message.channel)
                    if previous_content is None:
                        selected_file = await find_recent_supported_attachment(message.channel)
                        if selected_file is not None:
                            answer = await analyze_selected_attachment(
                                selected_file,
                                "Co si o této příloze myslíš?",
                            )
                            await message.reply(answer, mention_author=False)
                            return

                    if previous_content is None:
                        await message.channel.send(
                            "Nevím, k čemu se mám vyjádřit. Pošli text, přílohu, nebo to napiš pod zprávu."
                        )
                        return

                    await send_channel_chunks(
                        message,
                        await ask_panam(
                            OPENAI_MODEL,
                            "Co si o tom myslíš?\n\n" + previous_content,
                        ),
                    )
                    return

                await send_channel_chunks(message, await ask_panam(OPENAI_MODEL, value))
                return

            if intent_name == "ask_previous":
                selected_file = find_supported_attachment_in_message(message)
                if selected_file is not None:
                    answer = await analyze_selected_attachment(
                        selected_file,
                        "Co si o této příloze myslíš?",
                    )
                    await message.reply(answer, mention_author=False)
                    return

                previous_content = await find_recent_text_message(message.channel)
                if previous_content is None:
                    selected_file = await find_recent_supported_attachment(message.channel)
                    if selected_file is not None:
                        answer = await analyze_selected_attachment(
                            selected_file,
                            "Co si o této příloze myslíš?",
                        )
                        await message.reply(answer, mention_author=False)
                        return

                if previous_content is None:
                    await message.channel.send(
                        "Nevím, k čemu se mám vyjádřit. Pošli text, přílohu, nebo to napiš pod zprávu."
                    )
                    return

                await send_channel_chunks(
                    message,
                    await ask_panam(
                        OPENAI_MODEL,
                        "Co si o tom myslíš?\n\n" + previous_content,
                    ),
                )
                return

            if intent_name == "summary" and value:
                inline_summary = extract_inline_content_after_trigger(
                    request_text,
                    list(SUMMARY_TRIGGERS),
                )
                if inline_summary:
                    value = inline_summary

                selected_file = find_supported_attachment_in_message(message)
                if selected_file is not None:
                    answer = await analyze_selected_attachment(
                        selected_file,
                        "Shrň tuto přílohu.",
                    )
                    await message.reply(answer, mention_author=False)
                    return

                if is_context_reference(value):
                    previous_content = await find_recent_text_message(message.channel)
                    if previous_content is None:
                        await message.channel.send(
                            "Nemám co shrnout. Pošli text, přílohu, nebo to napiš hned pod zprávu, kterou chceš shrnout."
                        )
                        return

                    await send_channel_chunks(
                        message,
                        await summarize_text(OPENAI_MODEL, previous_content),
                    )
                    return

                await send_channel_chunks(message, await summarize_text(OPENAI_MODEL, value))
                return

            if intent_name == "summary_previous":
                selected_file = find_supported_attachment_in_message(message)
                if selected_file is not None:
                    answer = await analyze_selected_attachment(
                        selected_file,
                        "Shrň tuto přílohu.",
                    )
                    await message.reply(answer, mention_author=False)
                    return

                previous_content = await find_recent_text_message(message.channel)
                if previous_content is None:
                    await message.channel.send(
                        "Nemám co shrnout. Pošli text, přílohu, nebo to napiš hned pod zprávu, kterou chceš shrnout."
                    )
                    return

                await send_channel_chunks(
                    message,
                    await summarize_text(OPENAI_MODEL, previous_content),
                )
                return

            if intent_name == "talk" and value:
                await send_channel_chunks(
                    message,
                    await ask_panam_talk(OPENAI_MODEL, value),
                )
                return

            basic_prompt = extract_basic_panam_prompt(message.content or "")
            if basic_prompt is not None:
                await handle_basic_panam_message(message, basic_prompt)
                return

        except Exception:
            natural_status = "error"
            logger.exception("Chyba při zpracování přirozené zprávy")
            await message.channel.send("Něco se pokazilo při zpracování akce.")
        finally:
            log_action("natural_message", natural_action, natural_status, message)


bot = DiscordAIBot()


@bot.tree.command(
    name="ping",
    description="Ověř, že je bot online."
)
@log_slash_command("ping")
async def ping(interaction: discord.Interaction) -> None:
    if not is_interaction_allowed(interaction, "ping"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    await interaction.response.send_message("Panam je online.")


@bot.tree.command(
    name="help",
    description="Zobraz dostupné commandy."
)
@log_slash_command("help")
async def help_command(interaction: discord.Interaction) -> None:
    if not is_interaction_allowed(interaction, "help"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    await interaction.response.send_message(get_help_text())


@bot.tree.command(
    name="memory_clear",
    description="Vymaž krátkou konverzační paměť Panam pro tento kanál."
)
@log_slash_command("memory_clear")
async def memory_clear(interaction: discord.Interaction) -> None:
    channel_id = interaction.channel_id
    if channel_id is None:
        await interaction.response.send_message(
            "Tenhle command potřebuje běžet v kanálu.",
            ephemeral=True,
        )
        return

    if not is_interaction_allowed(interaction, "memory_clear"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    panam_memory.clear_channel_memory(channel_id)
    clear_last_file_context(channel_id)
    clear_last_router_decision(channel_id)
    await interaction.response.send_message(
        "Krátká paměť pro tento kanál je vymazaná."
    )


@bot.tree.command(
    name="note_add",
    description="Ulož krátkou poznámku."
)
@app_commands.describe(text="Text poznámky")
@log_slash_command("note_add")
async def note_add(interaction: discord.Interaction, text: str) -> None:
    if not is_interaction_allowed(interaction, "note_add"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    try:
        notes = load_notes()
        notes.append(
            {
                "text": text,
                "author_id": interaction.user.id,
                "author_name": getattr(
                    interaction.user,
                    "display_name",
                    interaction.user.name,
                ),
                "channel_id": interaction.channel_id,
                "created_at": datetime.now(timezone.utc).isoformat(),
            }
        )
        save_notes(notes)

        await interaction.response.send_message("Poznámka uložená.")

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /note_add")
        await interaction.response.send_message(
            "Něco se pokazilo při ukládání poznámky."
        )


@bot.tree.command(
    name="note_list",
    description="Vypiš poslední poznámky."
)
@log_slash_command("note_list")
async def note_list(interaction: discord.Interaction) -> None:
    if not is_interaction_allowed(interaction, "note_list"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    try:
        notes = load_notes()

        if not notes:
            await interaction.response.send_message("Zatím nemám žádné poznámky.")
            return

        lines = ["Poslední poznámky:"]
        for index, note in enumerate(reversed(notes[-10:]), start=1):
            author_name = note.get("author_name", "neznámý autor")
            created_at = note.get("created_at", "neznámý čas")
            note_text = note.get("text", "")
            lines.append(f"{index}. [{created_at}] {author_name}: {note_text}")

        answer = "\n".join(lines)
        answer = shorten_for_discord(answer)

        await interaction.response.send_message(answer)

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /note_list")
        await interaction.response.send_message(
            "Něco se pokazilo při načítání poznámek."
        )


@bot.tree.command(
    name="note_search",
    description="Vyhledej uložené poznámky."
)
@app_commands.describe(query="Text k vyhledání")
@log_slash_command("note_search")
async def note_search(interaction: discord.Interaction, query: str) -> None:
    if not is_interaction_allowed(interaction, "note_search"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    try:
        notes = load_notes()

        if not notes:
            await interaction.response.send_message("Zatím nemám žádné poznámky.")
            return

        query_lower = query.lower()
        matches = [
            note
            for note in notes
            if query_lower in str(note.get("text", "")).lower()
        ][-10:]

        if not matches:
            await interaction.response.send_message("Nic jsem nenašla.")
            return

        lines = ["Nalezené poznámky:"]
        for index, note in enumerate(reversed(matches), start=1):
            author_name = note.get("author_name", "neznámý autor")
            created_at = note.get("created_at", "neznámý čas")
            note_text = note.get("text", "")
            lines.append(f"{index}. [{created_at}] {author_name}: {note_text}")

        answer = "\n".join(lines)
        answer = shorten_for_discord(answer)

        await interaction.response.send_message(answer)

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /note_search")
        await interaction.response.send_message(
            "Něco se pokazilo při vyhledávání poznámek."
        )


@bot.tree.command(
    name="todo_add",
    description="Přidej úkol do todo listu."
)
@app_commands.describe(text="Text úkolu")
@log_slash_command("todo_add")
async def todo_add(interaction: discord.Interaction, text: str) -> None:
    if not is_interaction_allowed(interaction, "todo_add"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    try:
        todos = load_todos()
        next_id = max(
            (todo.get("id", 0) for todo in todos if isinstance(todo.get("id"), int)),
            default=0,
        ) + 1

        todos.append(
            {
                "id": next_id,
                "text": text,
                "done": False,
                "author_id": interaction.user.id,
                "author_name": getattr(
                    interaction.user,
                    "display_name",
                    interaction.user.name,
                ),
                "channel_id": interaction.channel_id,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "completed_at": None,
            }
        )
        save_todos(todos)

        await interaction.response.send_message(f"Úkol #{next_id} uložený.")

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /todo_add")
        await interaction.response.send_message(
            "Něco se pokazilo při ukládání úkolu."
        )


@bot.tree.command(
    name="todo_list",
    description="Vypiš aktivní úkoly."
)
@log_slash_command("todo_list")
async def todo_list(interaction: discord.Interaction) -> None:
    if not is_interaction_allowed(interaction, "todo_list"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    try:
        todos = load_todos()
        active_todos = [todo for todo in todos if not todo.get("done")]

        if not active_todos:
            await interaction.response.send_message("Nemáš žádné aktivní úkoly.")
            return

        lines = ["Aktivní úkoly:"]
        for todo in active_todos[:15]:
            todo_id = todo.get("id", "?")
            text = todo.get("text", "")
            author_name = todo.get("author_name", "neznámý autor")
            lines.append(f"#{todo_id} - {text} ({author_name})")

        answer = "\n".join(lines)
        answer = shorten_for_discord(answer)

        await interaction.response.send_message(answer)

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /todo_list")
        await interaction.response.send_message(
            "Něco se pokazilo při načítání úkolů."
        )


@bot.tree.command(
    name="todo_done",
    description="Označ úkol jako hotový."
)
@app_commands.describe(todo_id="ID úkolu")
@log_slash_command("todo_done")
async def todo_done(interaction: discord.Interaction, todo_id: int) -> None:
    if not is_interaction_allowed(interaction, "todo_done"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    try:
        todos = load_todos()
        todo = next((item for item in todos if item.get("id") == todo_id), None)

        if todo is None:
            await interaction.response.send_message("Takový úkol jsem nenašla.")
            return

        todo["done"] = True
        todo["completed_at"] = datetime.now(timezone.utc).isoformat()
        save_todos(todos)

        await interaction.response.send_message(f"Úkol #{todo_id} je hotový.")

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /todo_done")
        await interaction.response.send_message(
            "Něco se pokazilo při dokončování úkolu."
        )


@bot.tree.command(
    name="search_messages",
    description="Vyhledej text v posledních zprávách kanálu."
)
@app_commands.describe(
    query="Text k vyhledání",
    limit="Kolik posledních zpráv prohledat, maximálně 300",
)
@log_slash_command("search_messages")
async def search_messages(
    interaction: discord.Interaction,
    query: str,
    limit: int = 100,
) -> None:
    if not is_interaction_allowed(interaction, "search_messages"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    await interaction.response.defer(thinking=True)

    try:
        channel = interaction.channel
        history = getattr(channel, "history", None)
        if history is None:
            await interaction.followup.send(
                "Něco se pokazilo při vyhledávání zpráv."
            )
            return

        query_lower = query.lower()
        search_limit = min(max(limit, 1), 300)
        matches = []

        async for message in history(limit=search_limit):
            if message.author.bot:
                continue

            content = message.content or ""
            content_lower = content.lower()
            if query_lower not in content_lower:
                continue

            match_index = content_lower.find(query_lower)
            start = max(match_index - 45, 0)
            end = min(match_index + len(query) + 90, len(content))
            excerpt = content[start:end].replace("\n", " ").strip()

            if start > 0:
                excerpt = "…" + excerpt
            if end < len(content):
                excerpt = excerpt + "…"

            created_at = message.created_at.astimezone(timezone.utc).strftime(
                "%Y-%m-%d %H:%M UTC"
            )
            author_name = getattr(message.author, "display_name", message.author.name)
            matches.append(
                f"{len(matches) + 1}. [{created_at}] {author_name}: {excerpt}"
            )

            if len(matches) >= 10:
                break

        if not matches:
            await interaction.followup.send("Nic jsem nenašla.")
            return

        answer = "Nalezené zprávy:\n" + "\n".join(matches)
        answer = shorten_for_discord(answer)

        await interaction.followup.send(answer)

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /search_messages")
        await interaction.followup.send(
            "Něco se pokazilo při vyhledávání zpráv."
        )


@bot.tree.command(
    name="channel_summary",
    description="Shrň poslední zprávy z aktuálního kanálu."
)
@app_commands.describe(limit="Kolik posledních zpráv shrnout, maximálně 200")
@log_slash_command("channel_summary")
async def channel_summary(
    interaction: discord.Interaction,
    limit: int = 50,
) -> None:
    if not is_interaction_allowed(interaction, "channel_summary"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    await interaction.response.defer(thinking=True)

    try:
        channel = interaction.channel
        history = getattr(channel, "history", None)
        if history is None:
            await interaction.followup.send(
                "Něco se pokazilo při načítání zpráv."
            )
            return

        summary_limit = min(max(limit, 1), 200)
        messages = []

        async for message in history(limit=summary_limit):
            if message.author.bot:
                continue

            content = (message.content or "").strip()
            if not content:
                continue

            created_at = message.created_at.astimezone(timezone.utc).strftime(
                "%Y-%m-%d %H:%M UTC"
            )
            author_name = getattr(message.author, "display_name", message.author.name)
            messages.append(
                {
                    "author": author_name,
                    "created_at": created_at,
                    "content": content,
                }
            )

        if not messages:
            await interaction.followup.send("Nemám tu co shrnout.")
            return

        messages.reverse()
        channel_text = "\n".join(
            (
                f"Autor: {message['author']}\n"
                f"Čas: {message['created_at']}\n"
                f"Text: {message['content']}"
            )
            for message in messages
        )

        answer = await summarize_channel_messages(OPENAI_MODEL, channel_text)
        await interaction.followup.send(answer)

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /channel_summary")
        await interaction.followup.send(
            "Něco se pokazilo při shrnování kanálu."
        )


@bot.tree.command(
    name="summary",
    description="Stručně shrň delší text."
)
@app_commands.describe(text="Text ke shrnutí")
@log_slash_command("summary")
async def summary(interaction: discord.Interaction, text: str) -> None:
    if not is_interaction_allowed(interaction, "summary"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    await interaction.response.defer(thinking=True)

    try:
        answer = await summarize_text(OPENAI_MODEL, text)
        await interaction.followup.send(answer)

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /summary")
        await interaction.followup.send(
            "Něco se pokazilo při shrnování textu. Mrkni do konzole na chybu."
        )


@bot.tree.command(
    name="analyze",
    description="Analyzuj přiložený obrázek nebo dokument pomocí AI."
)
@app_commands.describe(
    file="Soubor k analýze",
    question="Co chceš k souboru zjistit",
)
@log_slash_command("analyze")
async def analyze(
    interaction: discord.Interaction,
    file: discord.Attachment | None = None,
    question: str = "Popiš, co je v příloze.",
) -> None:
    if not is_interaction_allowed(interaction, "analyze"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    await interaction.response.defer(thinking=True)

    selected_file = file
    if selected_file is None:
        selected_file = await find_recent_supported_attachment(interaction.channel)

    if selected_file is None:
        await interaction.followup.send(
            "Nevidím žádnou podporovanou přílohu ani v commandu, ani v předchozí zprávě."
        )
        return

    log_attachment_info("slash_command", "analyze", interaction, selected_file)

    try:
        answer = await analyze_selected_attachment(selected_file, question)
        await interaction.followup.send(answer)
        set_last_file_context(
            interaction.channel_id,
            Path(selected_file.filename).name,
            get_file_extension(selected_file.filename),
            "chat_answer",
            file_summary=answer,
        )

    except AttachmentAnalysisUserError as error:
        await interaction.followup.send(str(error))

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /analyze")
        await interaction.followup.send(
            "Něco se pokazilo při analýze přílohy. Mrkni do konzole na chybu."
        )


@bot.tree.command(
    name="read_file",
    description="Přečti TXT, MD, CSV, PDF, DOCX nebo XLSX přílohu pomocí AI."
)
@app_commands.describe(
    file="Dokument k přečtení",
    question="Co chceš k dokumentu zjistit",
)
@log_slash_command("read_file")
async def read_file(
    interaction: discord.Interaction,
    file: discord.Attachment,
    question: str = "Shrň mi tento dokument.",
) -> None:
    if not is_interaction_allowed(interaction, "read_file"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    if file.size > MAX_DOCUMENT_SIZE_BYTES:
        await interaction.response.send_message(
            "Ten soubor je moc velký. Zatím beru max 20 MB.",
            ephemeral=True,
        )
        return

    extension = get_file_extension(file.filename)
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        await interaction.response.send_message(
            "Tenhle typ dokumentu zatím neumím přečíst. Pošli mi prosím TXT, MD, CSV, PDF, DOCX nebo XLSX.",
            ephemeral=True,
        )
        return

    log_attachment_info("slash_command", "read_file", interaction, file)

    await interaction.response.defer(thinking=True)

    try:
        data = await read_attachment_bytes(file)

        if extension == ".pdf":
            document_text = extract_text_from_pdf(data)
            if not document_text:
                await interaction.followup.send(
                    "Z toho PDF se mi nepodařilo vytáhnout žádný text. Možná je to sken nebo obrázkové PDF."
                )
                return
        elif extension == ".docx":
            document_text = extract_text_from_docx(data)
        elif extension == ".xlsx":
            document_text = extract_text_from_xlsx(data)
        else:
            document_text = extract_text_from_plain_file(data)

        if not document_text.strip():
            if extension == ".xlsx":
                await interaction.followup.send(
                    "Z toho Excelu se mi nepodarilo vytahnout zadna data."
                )
                return

            await interaction.followup.send(
                "Z toho dokumentu se mi nepodařilo vytáhnout žádný text."
            )
            return

        document_text = trim_document_text(document_text)
        answer = await analyze_document_text(
            OPENAI_MODEL,
            document_text,
            question,
            file.filename,
        )
        await interaction.followup.send(answer)
        set_last_file_context(
            interaction.channel_id,
            Path(file.filename).name,
            extension,
            "chat_answer",
            file_summary=answer,
        )

    except AttachmentAnalysisUserError as error:
        mark_command_status(interaction, "error")
        await interaction.followup.send(str(error))

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /read_file")
        await interaction.followup.send(
            "Něco se pokazilo při čtení dokumentu. Mrkni do konzole na chybu."
        )


@bot.tree.command(
    name="file_job_test",
    description="Otestuj docasnou file-job pipeline na dokumentove priloze."
)
@app_commands.describe(
    file="Dokument k otestovani pipeline",
    output_format="md nebo txt",
)
@log_slash_command("file_job_test")
async def file_job_test(
    interaction: discord.Interaction,
    file: discord.Attachment,
    output_format: str = "md",
) -> None:
    if not is_interaction_allowed(interaction, "file_job_test"):
        await interaction.response.send_message(
            "Tady nemam povolene odpovidat.",
            ephemeral=True,
        )
        return

    if file.size > MAX_DOCUMENT_SIZE_BYTES:
        await interaction.response.send_message(
            "Ten soubor je moc velky. Zatim beru max 20 MB.",
            ephemeral=True,
        )
        return

    extension = get_file_extension(file.filename)
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        await interaction.response.send_message(
            "File job test zatim podporuje dokumenty TXT, MD, CSV, PDF, DOCX nebo XLSX.",
            ephemeral=True,
        )
        return

    output_extension = ".md" if output_format.lower().strip(".") == "md" else ".txt"
    job = None
    await interaction.response.defer(thinking=True)

    try:
        job = panam_files.create_file_job(
            user_id=interaction.user.id,
            channel_id=interaction.channel_id or 0,
            action="file_job_test",
        )
        log_action(
            "slash_command",
            "file_job_test",
            "started",
            interaction,
            job_id=job.job_id,
            filename=Path(file.filename).name,
            extension=extension,
            size=file.size,
        )

        input_path = await panam_files.save_attachment_to_job(file, job)
        data = input_path.read_bytes()
        extracted_text = extract_text_from_attachment(input_path.name, data)
        if not extracted_text.strip():
            extracted_text = "Z dokumentu se nepodarilo vytahnout zadny text."
        else:
            extracted_text = trim_document_text(extracted_text)
        panam_files.write_work_text(job, "extracted_text.txt", extracted_text)

        output_content = (
            "# Panam file job test\n\n"
            f"- job_id: `{job.job_id}`\n"
            f"- input: `{input_path.name}`\n\n"
            "## Extracted text\n\n"
            f"{extracted_text}\n"
        )
        if output_extension == ".txt":
            output_content = (
                "Panam file job test\n\n"
                f"job_id: {job.job_id}\n"
                f"input: {input_path.name}\n\n"
                "Extracted text\n\n"
                f"{extracted_text}\n"
            )

        output_path = panam_files.write_output_text(
            job,
            f"file_job_test_output{output_extension}",
            output_content,
        )
        job.status = "success"
        job.finished_at = datetime.now(timezone.utc).isoformat()
        panam_files.write_job_metadata(job)

        await interaction.followup.send(
            "File job hotovy.",
            file=discord.File(output_path),
        )
        log_action(
            "slash_command",
            "file_job_test",
            "success",
            interaction,
            job_id=job.job_id,
            filename=Path(file.filename).name,
            extension=extension,
            size=file.size,
        )

    except Exception:
        mark_command_status(interaction, "error")
        if job is not None:
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                "slash_command",
                "file_job_test",
                "error",
                interaction,
                job_id=job.job_id,
                filename=Path(file.filename).name,
                extension=extension,
                size=file.size,
            )
        logger.exception("Chyba pri zpracovani /file_job_test")
        await interaction.followup.send(
            "Neco se pokazilo pri testu file-job pipeline."
        )

    finally:
        if job is not None:
            panam_files.cleanup_job(job)


@bot.tree.command(
    name="transform_excel",
    description="Bezpecne uprav XLSX podle jednoduche instrukce a vrat novy soubor."
)
@app_commands.describe(
    file="XLSX soubor k uprave",
    instruction="Jednoducha deterministicka uprava Excelu",
)
@log_slash_command("transform_excel")
async def transform_excel(
    interaction: discord.Interaction,
    file: discord.Attachment,
    instruction: str,
) -> None:
    if not is_interaction_allowed(interaction, "transform_excel"):
        await interaction.response.send_message(
            "Tady nemam povolene odpovidat.",
            ephemeral=True,
        )
        return

    if file.size > MAX_DOCUMENT_SIZE_BYTES:
        await interaction.response.send_message(
            "Ten soubor je moc velky. Zatim beru max 20 MB.",
            ephemeral=True,
        )
        return

    extension = get_file_extension(file.filename)
    if extension != ".xlsx":
        await interaction.response.send_message(
            "Tenhle command podporuje jen XLSX soubory.",
            ephemeral=True,
        )
        return

    await interaction.response.defer(thinking=True)
    await run_spreadsheet_transform_file_job(
        interaction,
        file,
        instruction,
        "slash_command",
    )


@bot.tree.command(
    name="process_file",
    description="Zpracuj dokument podle instrukce a vrat vystup jako soubor."
)
@app_commands.describe(
    file="Dokument ke zpracovani",
    instruction="Co ma Panam s dokumentem udelat",
    output_format="md, txt nebo docx",
)
@log_slash_command("process_file")
async def process_file(
    interaction: discord.Interaction,
    file: discord.Attachment,
    instruction: str,
    output_format: str = "md",
) -> None:
    if not is_interaction_allowed(interaction, "process_file"):
        await interaction.response.send_message(
            "Tady nemam povolene odpovidat.",
            ephemeral=True,
        )
        return

    if file.size > MAX_DOCUMENT_SIZE_BYTES:
        await interaction.response.send_message(
            "Ten soubor je moc velky. Zatim beru max 20 MB.",
            ephemeral=True,
        )
        return

    extension = get_file_extension(file.filename)
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        await interaction.response.send_message(
            "Tenhle command podporuje dokumenty TXT, MD, CSV, PDF, DOCX nebo XLSX.",
            ephemeral=True,
        )
        return

    normalized_output_format = output_format.lower().strip(".")
    if normalized_output_format not in {"md", "txt", "docx"}:
        await interaction.response.send_message(
            "Podporovane output_format jsou jen md, txt nebo docx.",
            ephemeral=True,
        )
        return

    output_extension = f".{normalized_output_format}"
    job = None
    await interaction.response.defer(thinking=True)

    try:
        job = panam_files.create_file_job(
            user_id=interaction.user.id,
            channel_id=interaction.channel_id or 0,
            action="process_file",
        )
        log_action(
            "slash_command",
            "process_file",
            "started",
            interaction,
            job_id=job.job_id,
            filename=Path(file.filename).name,
            extension=extension,
            size_bytes=file.size,
            output_format=normalized_output_format,
        )

        input_path = await panam_files.save_attachment_to_job(file, job)
        log_action(
            "slash_command",
            "process_file",
            "attachment_saved",
            interaction,
            job_id=job.job_id,
            filename=input_path.name,
            extension=input_path.suffix.lower(),
            size_bytes=input_path.stat().st_size,
            output_format=normalized_output_format,
        )

        data = input_path.read_bytes()
        extracted_text = extract_text_from_attachment(input_path.name, data)
        if not extracted_text.strip():
            mark_command_status(interaction, "error")
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                "slash_command",
                "process_file",
                "error",
                interaction,
                job_id=job.job_id,
                filename=input_path.name,
                extension=input_path.suffix.lower(),
                size_bytes=input_path.stat().st_size,
                output_format=normalized_output_format,
            )
            message_text = (
                "Z toho Excelu se mi nepodarilo vytahnout zadna data."
                if extension == ".xlsx"
                else "Z toho dokumentu se mi nepodarilo vytahnout zadny text."
            )
            await interaction.followup.send(message_text)
            return

        extracted_text = trim_document_text(extracted_text)
        work_path = panam_files.write_work_text(job, "extracted_text.txt", extracted_text)
        log_action(
            "slash_command",
            "process_file",
            "text_extracted",
            interaction,
            job_id=job.job_id,
            filename=work_path.name,
            extension=work_path.suffix.lower(),
            size_bytes=work_path.stat().st_size,
            output_format=normalized_output_format,
        )

        process_output_format = (
            "md" if normalized_output_format == "docx" else normalized_output_format
        )
        processed_text = await process_document_text(
            OPENAI_MODEL,
            extracted_text,
            instruction,
            input_path.name,
            process_output_format,
        )
        output_filename = panam_files.build_panam_output_filename(
            input_path.name,
            output_extension,
        )
        if normalized_output_format == "docx":
            output_path = panam_files.ensure_within_job(
                job.output_dir / output_filename,
                job,
            )
            panam_docx.create_docx_from_text(processed_text, output_path)
            job.output_files.append(
                {
                    "filename": output_path.name,
                    "extension": output_path.suffix.lower(),
                    "size_bytes": output_path.stat().st_size,
                }
            )
            panam_files.write_job_metadata(job)
        else:
            output_path = panam_files.write_output_text(
                job,
                output_filename,
                processed_text,
            )
        log_action(
            "slash_command",
            "process_file",
            "output_written",
            interaction,
            job_id=job.job_id,
            filename=output_path.name,
            extension=output_path.suffix.lower(),
            size_bytes=output_path.stat().st_size,
            output_format=normalized_output_format,
        )

        job.status = "success"
        job.finished_at = datetime.now(timezone.utc).isoformat()
        panam_files.write_job_metadata(job)

        await interaction.followup.send(
            "Soubor je zpracovany.",
            file=discord.File(output_path),
        )
        set_last_file_context(
            interaction.channel_id,
            Path(file.filename).name,
            extension,
            "human_document",
            output_path.name,
        )
        log_action(
            "slash_command",
            "process_file",
            "output_sent",
            interaction,
            job_id=job.job_id,
            filename=output_path.name,
            extension=output_path.suffix.lower(),
            size_bytes=output_path.stat().st_size,
            output_format=normalized_output_format,
        )
        log_action(
            "slash_command",
            "process_file",
            "success",
            interaction,
            job_id=job.job_id,
            filename=Path(file.filename).name,
            extension=extension,
            size_bytes=file.size,
            output_format=normalized_output_format,
        )


    except AttachmentAnalysisUserError as error:
        mark_command_status(interaction, "error")
        if job is not None:
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                "slash_command",
                "process_file",
                "error",
                interaction,
                job_id=job.job_id,
                filename=Path(file.filename).name,
                extension=extension,
                size_bytes=file.size,
                output_format=normalized_output_format,
            )
        logger.exception("Chyba pri zpracovani /process_file")
        await interaction.followup.send(str(error))

    except Exception:
        mark_command_status(interaction, "error")
        if job is not None:
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                "slash_command",
                "process_file",
                "error",
                interaction,
                job_id=job.job_id,
                filename=Path(file.filename).name,
                extension=extension,
                size_bytes=file.size,
                output_format=normalized_output_format,
            )
        logger.exception("Chyba pri zpracovani /process_file")
        await interaction.followup.send(
            "Neco se pokazilo pri zpracovani souboru."
        )

    finally:
        if job is not None:
            panam_files.cleanup_job(job)


@bot.tree.command(
    name="extract_data",
    description="Vytez ze souboru strukturovana data jako JSON, CSV, Markdown nebo XLSX."
)
@app_commands.describe(
    file="Dokument pro tezeni dat",
    instruction="Jaka data ma Panam vytahnout",
    output_format="json, csv, md nebo xlsx",
)
@log_slash_command("extract_data")
async def extract_data(
    interaction: discord.Interaction,
    file: discord.Attachment,
    instruction: str,
    output_format: str = "json",
) -> None:
    if not is_interaction_allowed(interaction, "extract_data"):
        await interaction.response.send_message(
            "Tady nemam povolene odpovidat.",
            ephemeral=True,
        )
        return

    if file.size > MAX_DOCUMENT_SIZE_BYTES:
        await interaction.response.send_message(
            "Ten soubor je moc velky. Zatim beru max 20 MB.",
            ephemeral=True,
        )
        return

    extension = get_file_extension(file.filename)
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        await interaction.response.send_message(
            "Tenhle command podporuje dokumenty TXT, MD, CSV, PDF, DOCX nebo XLSX.",
            ephemeral=True,
        )
        return

    normalized_output_format = output_format.lower().strip(".")
    if normalized_output_format not in {"json", "csv", "md", "xlsx"}:
        await interaction.response.send_message(
            "Podporovane output_format jsou jen json, csv, md nebo xlsx.",
            ephemeral=True,
        )
        return

    output_extension = f".{normalized_output_format}"
    job = None
    await interaction.response.defer(thinking=True)

    try:
        job = panam_files.create_file_job(
            user_id=interaction.user.id,
            channel_id=interaction.channel_id or 0,
            action="extract_data",
        )
        log_action(
            "slash_command",
            "extract_data",
            "started",
            interaction,
            job_id=job.job_id,
            filename=Path(file.filename).name,
            extension=extension,
            size_bytes=file.size,
            output_format=normalized_output_format,
        )

        input_path = await panam_files.save_attachment_to_job(file, job)
        log_action(
            "slash_command",
            "extract_data",
            "attachment_saved",
            interaction,
            job_id=job.job_id,
            filename=input_path.name,
            extension=input_path.suffix.lower(),
            size_bytes=input_path.stat().st_size,
            output_format=normalized_output_format,
        )

        data = input_path.read_bytes()
        extracted_text = extract_text_from_attachment(input_path.name, data)
        if not extracted_text.strip():
            mark_command_status(interaction, "error")
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                "slash_command",
                "extract_data",
                "error",
                interaction,
                job_id=job.job_id,
                filename=input_path.name,
                extension=input_path.suffix.lower(),
                size_bytes=input_path.stat().st_size,
                output_format=normalized_output_format,
            )
            message_text = (
                "Z toho Excelu se mi nepodarilo vytahnout zadna data."
                if extension == ".xlsx"
                else "Z toho dokumentu se mi nepodarilo vytahnout zadny text."
            )
            await interaction.followup.send(message_text)
            return

        extracted_text = trim_document_text(extracted_text)
        work_path = panam_files.write_work_text(job, "extracted_text.txt", extracted_text)
        log_action(
            "slash_command",
            "extract_data",
            "text_extracted",
            interaction,
            job_id=job.job_id,
            filename=work_path.name,
            extension=work_path.suffix.lower(),
            size_bytes=work_path.stat().st_size,
            output_format=normalized_output_format,
        )

        structured_text = await extract_structured_data(
            OPENAI_MODEL,
            extracted_text,
            instruction,
            input_path.name,
            "json" if normalized_output_format == "xlsx" else normalized_output_format,
        )
        log_action(
            "slash_command",
            "extract_data",
            "ai_extraction_completed",
            interaction,
            job_id=job.job_id,
            filename=input_path.name,
            extension=input_path.suffix.lower(),
            size_bytes=input_path.stat().st_size,
            output_format=normalized_output_format,
        )

        output_filename = panam_files.build_panam_output_filename(
            input_path.name,
            output_extension,
        )
        json_data = None
        if normalized_output_format in {"json", "xlsx"}:
            try:
                json_data = json.loads(structured_text)
            except json.JSONDecodeError:
                mark_command_status(interaction, "error")
                job.status = "error"
                job.finished_at = datetime.now(timezone.utc).isoformat()
                panam_files.write_job_metadata(job)
                log_action(
                    "slash_command",
                    "extract_data",
                    "error",
                    interaction,
                    job_id=job.job_id,
                    filename=input_path.name,
                    extension=input_path.suffix.lower(),
                    size_bytes=input_path.stat().st_size,
                    output_format=normalized_output_format,
                )
                logger.warning(
                    "extract_data JSON validation failed job_id=%s status=error",
                    job.job_id,
                )
                if normalized_output_format == "xlsx":
                    await interaction.followup.send(
                        "Data pro XLSX se nepodarilo validovat. Zkus presnejsi instrukci nebo output_format=json."
                    )
                else:
                    await interaction.followup.send(
                        "Vystup JSON se nepodarilo validovat. Zkus presnejsi instrukci nebo output_format=md."
                    )
                return
            if normalized_output_format == "xlsx":
                log_action(
                    "slash_command",
                    "extract_data",
                    "json_validated_for_xlsx",
                    interaction,
                    job_id=job.job_id,
                    filename=output_filename,
                    extension=".xlsx",
                    size_bytes=len(structured_text.encode("utf-8")),
                    output_format=normalized_output_format,
                )
        elif normalized_output_format == "csv":
            if not structured_text.strip() or not structured_text.strip().splitlines():
                raise ValueError("CSV output is empty.")
        elif not structured_text.strip():
            raise ValueError("Markdown output is empty.")

        log_action(
            "slash_command",
            "extract_data",
            "output_validated",
            interaction,
            job_id=job.job_id,
            filename=output_filename,
            extension=output_extension,
            size_bytes=len(structured_text.encode("utf-8")),
            output_format=normalized_output_format,
        )

        if normalized_output_format == "xlsx":
            output_path = panam_files.ensure_within_job(
                job.output_dir / output_filename,
                job,
            )
            panam_excel.create_xlsx_from_items(
                get_items_for_xlsx(json_data),
                output_path,
            )
            job.output_files.append(
                {
                    "filename": output_path.name,
                    "extension": output_path.suffix.lower(),
                    "size_bytes": output_path.stat().st_size,
                }
            )
            panam_files.write_job_metadata(job)
            output_status = "xlsx_output_written"
        else:
            output_path = panam_files.write_output_text(
                job,
                output_filename,
                structured_text,
            )
            output_status = "output_written"

        log_action(
            "slash_command",
            "extract_data",
            output_status,
            interaction,
            job_id=job.job_id,
            filename=output_path.name,
            extension=output_path.suffix.lower(),
            size_bytes=output_path.stat().st_size,
            output_format=normalized_output_format,
        )

        job.status = "success"
        job.finished_at = datetime.now(timezone.utc).isoformat()
        panam_files.write_job_metadata(job)

        await interaction.followup.send(
            "Data jsou vytezena.",
            file=discord.File(output_path),
        )
        set_last_file_context(
            interaction.channel_id,
            Path(file.filename).name,
            extension,
            "structured_data",
            output_path.name,
        )
        log_action(
            "slash_command",
            "extract_data",
            "output_sent",
            interaction,
            job_id=job.job_id,
            filename=output_path.name,
            extension=output_path.suffix.lower(),
            size_bytes=output_path.stat().st_size,
            output_format=normalized_output_format,
        )
        log_action(
            "slash_command",
            "extract_data",
            "success",
            interaction,
            job_id=job.job_id,
            filename=Path(file.filename).name,
            extension=extension,
            size_bytes=file.size,
            output_format=normalized_output_format,
        )


    except AttachmentAnalysisUserError as error:
        mark_command_status(interaction, "error")
        if job is not None:
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                "slash_command",
                "extract_data",
                "error",
                interaction,
                job_id=job.job_id,
                filename=Path(file.filename).name,
                extension=extension,
                size_bytes=file.size,
                output_format=normalized_output_format,
            )
        logger.exception("Chyba pri zpracovani /extract_data")
        await interaction.followup.send(str(error))

    except Exception:
        mark_command_status(interaction, "error")
        if job is not None:
            job.status = "error"
            job.finished_at = datetime.now(timezone.utc).isoformat()
            panam_files.write_job_metadata(job)
            log_action(
                "slash_command",
                "extract_data",
                "error",
                interaction,
                job_id=job.job_id,
                filename=Path(file.filename).name,
                extension=extension,
                size_bytes=file.size,
                output_format=normalized_output_format,
            )
        logger.exception("Chyba pri zpracovani /extract_data")
        await interaction.followup.send(
            "Neco se pokazilo pri tezeni dat ze souboru."
        )

    finally:
        if job is not None:
            panam_files.cleanup_job(job)


@bot.tree.command(
    name="ask",
    description="Pošli otázku do OpenAI API a vrať odpověď do Discordu."
)
@app_commands.describe(question="Tvoje otázka pro AI")
@log_slash_command("ask")
async def ask(interaction: discord.Interaction, question: str) -> None:
    if not is_interaction_allowed(interaction, "ask"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    await interaction.response.defer(thinking=True)

    try:
        answer = await ask_panam(OPENAI_MODEL, question)
        await interaction.followup.send(answer)

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /ask")
        await interaction.followup.send(
            "Něco se pokazilo při volání AI. Mrkni do konzole na chybu."
        )


@bot.tree.command(
    name="panam_talk",
    description="Promluv si s Panam v osobnějším talk režimu."
)
@app_commands.describe(message="Zpráva pro Panam talk režim")
@log_slash_command("panam_talk")
async def panam_talk(interaction: discord.Interaction, message: str) -> None:
    if not is_interaction_allowed(interaction, "panam_talk"):
        await interaction.response.send_message(
            "Tady nemám povolené odpovídat.",
            ephemeral=True,
        )
        return

    await interaction.response.defer(thinking=True)

    try:
        answer = await ask_panam_talk(OPENAI_MODEL, message)
        await send_followup_chunks(interaction, answer)

    except Exception:
        mark_command_status(interaction, "error")
        logger.exception("Chyba při zpracování /panam_talk")
        await interaction.followup.send(
            "Něco se pokazilo při talk režimu. Mrkni do konzole na chybu."
        )


if __name__ == "__main__":
    bot.run(DISCORD_BOT_TOKEN)
