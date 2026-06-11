import logging
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, cast

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


MIN_COLUMN_WIDTH = 10
MAX_COLUMN_WIDTH = 50
COLUMN_WIDTH_PADDING = 2
COLUMN_JOINER_WORDS = {"a", "and"}
HEADER_SCAN_LIMIT = 20

logger = logging.getLogger("panam.spreadsheet")


class SpreadsheetTransformUserError(Exception):
    pass


@dataclass(frozen=True)
class SpreadsheetTransformPlan:
    remove_empty_rows: bool = False
    selected_columns: tuple[str, ...] = ()
    filter_column: str | None = None
    filter_value: str | None = None
    filter_value_display: str | None = None
    sort_column: str | None = None
    duplicate_column: str | None = None


@dataclass(frozen=True)
class SpreadsheetTransformResult:
    output_path: Path
    operation_summary: str
    row_count_before: int | None
    row_count_after: int | None
    column_count_before: int | None
    column_count_after: int | None
    rows_read: int
    rows_written: int
    columns_written: int
    operations: tuple[str, ...]


SUPPORTED_TRANSFORM_EXAMPLES = (
    "odstran prazdne radky",
    "nech jen radky kde Stav = Chyba",
    "vyber sloupce Jmeno, Email",
    "serad podle Jmeno vzestupne",
    "najdi duplicity podle Email",
)


def unsupported_instruction_message() -> str:
    examples = "\n".join(f"- {example}" for example in SUPPORTED_TRANSFORM_EXAMPLES)
    return (
        "Excel transform v1 umi jen jednoduche deterministicke upravy. "
        "Neprovedla jsem zadnou AI editaci dat.\n\n"
        f"Podporovane priklady:\n{examples}"
    )


def _normalize_text(text: Any) -> str:
    normalized = unicodedata.normalize("NFKD", str(text or "").casefold())
    without_marks = "".join(
        char for char in normalized if not unicodedata.combining(char)
    )
    return re.sub(r"\s+", " ", without_marks).strip()


def normalize_column_name(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or "").casefold())
    without_marks = "".join(
        char for char in normalized if not unicodedata.combining(char)
    )
    without_punctuation = re.sub(r"[^\w\s]+", "", without_marks)
    return re.sub(r"\s+", " ", without_punctuation).strip()


def _column_tokens(value: str) -> list[str]:
    return [token for token in normalize_column_name(value).split() if token]


def _without_joiner_words(value: str) -> str:
    return " ".join(
        token for token in _column_tokens(value) if token not in COLUMN_JOINER_WORDS
    )


def _normalize_header(text: Any) -> str:
    return re.sub(r"\s+", "", normalize_column_name(str(text or "")))


def _ambiguous_column_error(requested_name: str, matches: list[str]) -> SpreadsheetTransformUserError:
    return SpreadsheetTransformUserError(
        f'Sloupec "{requested_name}" je nejednoznacny. Mozne sloupce: {", ".join(matches)}.'
    )


def find_column_index(headers: list[str], requested_name: str) -> int | None:
    requested = normalize_column_name(requested_name)
    if not requested:
        return None

    normalized_headers = [normalize_column_name(header) for header in headers]

    exact_matches = [
        index for index, header in enumerate(normalized_headers) if header == requested
    ]
    if len(exact_matches) == 1:
        return exact_matches[0]
    if len(exact_matches) > 1:
        raise _ambiguous_column_error(
            requested_name,
            [headers[index] for index in exact_matches],
        )

    requested_without_joiners = _without_joiner_words(requested)
    without_joiner_matches = [
        index
        for index, header in enumerate(headers)
        if _without_joiner_words(header) == requested_without_joiners
    ]
    if len(without_joiner_matches) == 1:
        return without_joiner_matches[0]
    if len(without_joiner_matches) > 1:
        raise _ambiguous_column_error(
            requested_name,
            [headers[index] for index in without_joiner_matches],
        )

    contains_matches = [
        index
        for index, header in enumerate(normalized_headers)
        if requested in header
    ]
    if len(contains_matches) == 1:
        return contains_matches[0]
    if len(contains_matches) > 1:
        raise _ambiguous_column_error(
            requested_name,
            [headers[index] for index in contains_matches],
        )

    requested_tokens = set(_column_tokens(requested))
    if requested_tokens:
        token_subset_matches = [
            index
            for index, header in enumerate(headers)
            if requested_tokens.issubset(set(_column_tokens(header)))
        ]
        if len(token_subset_matches) == 1:
            return token_subset_matches[0]
        if len(token_subset_matches) > 1:
            raise _ambiguous_column_error(
                requested_name,
                [headers[index] for index in token_subset_matches],
            )

    return None


def _is_empty_value(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _requested_columns_from_plan(plan: SpreadsheetTransformPlan) -> tuple[str, ...]:
    requested_columns: list[str] = []

    for column in (
        plan.filter_column,
        plan.sort_column,
        plan.duplicate_column,
        *plan.selected_columns,
    ):
        if column:
            requested_columns.append(column)

    return tuple(requested_columns)


def _headers_from_row(row: tuple[Any, ...]) -> list[str]:
    return [str(value).strip() if value is not None else "" for value in row]


def _has_unique_header_names(headers: list[str]) -> bool:
    normalized_seen: set[str] = set()

    for header in headers:
        normalized_header = _normalize_header(header)
        if not normalized_header:
            continue
        if normalized_header in normalized_seen:
            return False
        normalized_seen.add(normalized_header)

    return True


def _requested_column_match_count(headers: list[str], requested_columns: tuple[str, ...]) -> int:
    matches = 0

    for requested_column in requested_columns:
        try:
            if find_column_index(headers, requested_column) is not None:
                matches += 1
        except SpreadsheetTransformUserError:
            return -1

    return matches


def _choose_header_row(
    rows: list[tuple[Any, ...]],
    requested_columns: tuple[str, ...],
) -> tuple[int, list[str]]:
    best_score: tuple[int, int] | None = None
    best_candidate: tuple[int, list[str]] | None = None

    for row_index, row in enumerate(rows[:HEADER_SCAN_LIMIT]):
        headers = _headers_from_row(row)
        non_empty_headers = [header for header in headers if header]

        if len(non_empty_headers) < 2:
            continue

        if not _has_unique_header_names(headers):
            continue

        match_count = _requested_column_match_count(headers, requested_columns)
        if match_count < 0:
            continue

        score = (match_count, -row_index)
        if best_score is None or score > best_score:
            best_score = score
            best_candidate = (row_index, headers)

    if best_candidate is None:
        raise SpreadsheetTransformUserError(
            "V Excelu jsem nenašla řádek s hlavičkami sloupců."
        )

    return best_candidate


def _split_columns(text: str) -> tuple[str, ...]:
    clean_text = re.sub(r"\s+", " ", text).strip(" .,;:")
    if not clean_text:
        return ()

    parts = [
        part.strip(" .,;:")
        for part in re.split(r"\s*(?:,|;)\s*", clean_text)
        if part.strip(" .,;:")
    ]
    return tuple(parts)


def _extract_filter_display_value(instruction: str) -> str | None:
    match = re.search(
        r"\bkde\s+(.+?)\s*(?:=|\s+je\s+|\s+rovna\s+se\s+)\s*(.+?)(?=\s*(?:,|;|$|\bvyber\b|\bserad\b|\bseřaď\b|\bsetrid\b|\bsort\b|\bodstran\b|\bodstraň\b|\bsmaz\b|\bsmaž\b|\bnajdi\b|\bdetekuj\b|\bduplicity\b|\bduplicit\b))",
        str(instruction or ""),
        flags=re.IGNORECASE,
    )
    if not match:
        return None

    value = re.sub(r"\s+", " ", match.group(2)).strip(" .,;:")
    return value or None


def parse_transform_instruction(instruction: str) -> SpreadsheetTransformPlan:
    normalized = _normalize_text(instruction)
    if not normalized:
        raise SpreadsheetTransformUserError(unsupported_instruction_message())

    remove_empty_rows = bool(
        re.search(
            r"\b(odstran|odstrante|smaz|smazat|vyhod|odeber)\b.*\b(prazdne|empty)\b.*\b(radky|radku|rows)\b",
            normalized,
        )
    )

    selected_columns: tuple[str, ...] = ()
    select_match = re.search(
        r"\b(?:vyber|nech|ponech|zachovej)(?: jen)? sloupce?\s+(.+?)(?=\s+\b(?:kde|serad|setrid|sort|odstran|smaz|najdi|detekuj|duplicity|duplicit)\b|$)",
        normalized,
    )
    if select_match:
        selected_columns = _split_columns(select_match.group(1))

    filter_column = None
    filter_value = None
    filter_match = re.search(
        r"\bkde\s+([a-z0-9 _.-]+?)\s*(?:=| je | rovna se )\s*([a-z0-9 _.-]+?)(?=\s*(?:,|;|$|\bvyber\b|\bserad\b|\bsetrid\b|\bsort\b|\bodstran\b|\bsmaz\b|\bnajdi\b|\bdetekuj\b|\bduplicity\b|\bduplicit\b))",
        normalized,
    )
    if filter_match:
        filter_column = filter_match.group(1).strip()
        filter_value = filter_match.group(2).strip()

    sort_column = None
    sort_match = re.search(
        r"\b(?:serad|setrid|sort)\b(?:\s+\w+){0,3}?\s+podle\s+([a-z0-9 _.-]+?)(?=\s*(?:,|;|$|\bvzestupne\b|\basc\b|\bvyber\b|\bkde\b|\bodstran\b|\bnajdi\b|\bdetekuj\b))",
        normalized,
    )
    if sort_match and "sestupne" not in normalized and "desc" not in normalized:
        sort_column = sort_match.group(1).strip()

    duplicate_column = None
    duplicate_match = re.search(
        r"\b(?:duplicity|duplicit|duplikaty|duplikatu)\b.*\b(?:podle|ve sloupci|sloupec)\s+([a-z0-9 _.-]+?)(?=\s*(?:,|;|$|\bvyber\b|\bkde\b|\bserad\b|\bodstran\b))",
        normalized,
    )
    if duplicate_match:
        duplicate_column = duplicate_match.group(1).strip()

    plan = SpreadsheetTransformPlan(
        remove_empty_rows=remove_empty_rows,
        selected_columns=selected_columns,
        filter_column=filter_column,
        filter_value=filter_value,
        filter_value_display=_extract_filter_display_value(instruction),
        sort_column=sort_column,
        duplicate_column=duplicate_column,
    )
    if not any(
        (
            plan.remove_empty_rows,
            plan.selected_columns,
            plan.filter_column and plan.filter_value is not None,
            plan.sort_column,
            plan.duplicate_column,
        )
    ):
        raise SpreadsheetTransformUserError(unsupported_instruction_message())

    return plan


def _format_worksheet(worksheet: Worksheet) -> None:
    header_font = Font(bold=True)
    top_alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.freeze_panes = "A2"

    for cell in worksheet[1]:
        cell.font = header_font
        cell.alignment = top_alignment

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = top_alignment

    if worksheet.max_row and worksheet.max_column:
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
        )

    for column_index in range(1, worksheet.max_column + 1):
        max_length = 0
        for column_cells in worksheet.iter_cols(
            min_col=column_index,
            max_col=column_index,
            min_row=1,
            max_row=worksheet.max_row,
        ):
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))

        width = min(
            max(max_length + COLUMN_WIDTH_PADDING, MIN_COLUMN_WIDTH),
            MAX_COLUMN_WIDTH,
        )
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def _read_first_sheet_values(
    input_path: Path,
    requested_columns: tuple[str, ...] = (),
) -> tuple[list[str], list[list[Any]]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise SpreadsheetTransformUserError("Ten Excel je prazdny.")

    header_row_index, headers = _choose_header_row(rows, requested_columns)

    width = len(headers)
    data_rows = [
        list(row[:width]) + [None] * max(0, width - len(row))
        for row in rows[header_row_index + 1:]
    ]
    return headers, data_rows


def _header_index(headers: list[str], requested_column: str) -> int:
    column_index = find_column_index(headers, requested_column)
    if column_index is not None:
        return column_index

    raise SpreadsheetTransformUserError(
        f'Sloupec "{requested_column}" jsem v Excelu nenasla. '
        f'Dostupne sloupce: {", ".join(headers)}.'
    )


def _operation_names(plan: SpreadsheetTransformPlan) -> tuple[str, ...]:
    operations: list[str] = []
    if plan.remove_empty_rows:
        operations.append("remove_empty_rows")
    if plan.filter_column:
        operations.append("filter_equals")
    if plan.duplicate_column:
        operations.append("detect_duplicates")
    if plan.sort_column:
        operations.append("sort_ascending")
    if plan.selected_columns:
        operations.append("select_columns")
    return tuple(operations)


def _build_operation_summary(parts: list[str]) -> str:
    clean_parts = [part.strip() for part in parts if part.strip()]
    if not clean_parts:
        return "provedena podporovaná Excel transformace"
    return "; ".join(clean_parts)


def transform_xlsx(
    input_path: Path,
    output_path: Path,
    instruction: str,
    sheet_name: str = "Data",
) -> SpreadsheetTransformResult:
    plan = parse_transform_instruction(instruction)
    requested_columns = _requested_columns_from_plan(plan)
    headers, rows = _read_first_sheet_values(input_path, requested_columns)
    rows_read = len(rows)
    column_count_before = len(headers)
    operation_summary_parts: list[str] = []

    if plan.remove_empty_rows:
        rows = [row for row in rows if not all(_is_empty_value(value) for value in row)]
        operation_summary_parts.append("odstranění prázdných řádků")

    if plan.filter_column and plan.filter_value is not None:
        column_index = _header_index(headers, plan.filter_column)
        column_name = headers[column_index]
        display_value = plan.filter_value_display or plan.filter_value
        operation_summary_parts.append(
            f"filtr řádků, kde {column_name} = {display_value}"
        )
        wanted = _normalize_text(plan.filter_value)
        rows = [
            row
            for row in rows
            if _normalize_text(row[column_index] if column_index < len(row) else "") == wanted
        ]

    if plan.duplicate_column:
        column_index = _header_index(headers, plan.duplicate_column)
        column_name = headers[column_index]
        operation_summary_parts.append(
            f"nalezení duplicit podle sloupce {column_name}"
        )
        counts: dict[str, int] = {}
        keys: list[str] = []
        for row in rows:
            key = _normalize_text(row[column_index] if column_index < len(row) else "")
            keys.append(key)
            if key:
                counts[key] = counts.get(key, 0) + 1
        rows = [
            row
            for row, key in zip(rows, keys)
            if key and counts.get(key, 0) > 1
        ]

    if plan.sort_column:
        column_index = _header_index(headers, plan.sort_column)
        column_name = headers[column_index]
        operation_summary_parts.append(f"seřazení podle sloupce {column_name}")
        rows = sorted(
            rows,
            key=lambda row: _normalize_text(
                row[column_index] if column_index < len(row) else ""
            ),
        )

    selected_indexes = list(range(len(headers)))
    if plan.selected_columns:
        selected_indexes = [_header_index(headers, column) for column in plan.selected_columns]
        selected_headers = [headers[index] for index in selected_indexes]
        operation_summary_parts.append(
            f"výběr sloupců {', '.join(selected_headers)}"
        )
        headers = selected_headers
        rows = [[row[index] if index < len(row) else None for index in selected_indexes] for row in rows]

    workbook = Workbook()
    worksheet = cast(Worksheet, workbook.active)
    worksheet.title = sheet_name[:31] if sheet_name else "Data"
    worksheet.append(headers)
    for row in rows:
        worksheet.append(["" if value is None else value for value in row])

    _format_worksheet(worksheet)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

    operations = _operation_names(plan)
    operation_summary = _build_operation_summary(operation_summary_parts)
    logger.info(
        "spreadsheet_transform status=success input=%s output=%s rows_read=%s rows_written=%s columns_written=%s operations=%s",
        input_path.name,
        output_path.name,
        rows_read,
        len(rows),
        len(headers),
        ",".join(operations),
    )

    return SpreadsheetTransformResult(
        output_path=output_path,
        operation_summary=operation_summary,
        row_count_before=rows_read,
        row_count_after=len(rows),
        column_count_before=column_count_before,
        column_count_after=len(headers),
        rows_read=rows_read,
        rows_written=len(rows),
        columns_written=len(headers),
        operations=operations,
    )
